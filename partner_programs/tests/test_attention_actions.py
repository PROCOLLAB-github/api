import io
from concurrent.futures import ThreadPoolExecutor
from unittest import skipUnless
from unittest.mock import patch

from django.core import mail, signing
from django.db import close_old_connections, connection, transaction
from django.test import TransactionTestCase, override_settings
from openpyxl import load_workbook

from partner_programs.models import (
    ProgramAttentionReminder,
    ProgramAttentionReminderBatch,
)
from partner_programs.services.attention_actions import (
    confirm_reminder,
    deliver_reminder,
    dispatch_pending_reminders,
    preview_reminder,
)
from partner_programs.tests.test_attention_analytics_api import AttentionAnalyticsTestCase
from partner_programs.tests.helpers import (
    create_user,
    create_partner_program,
    create_program_member,
)
from project_rates.tests.helpers import create_rate_expert

VISITS = "participants-without-visit"
TEAMS = "participants-without-team"
WORKS = "projects-awaiting-evaluation"


@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
class AttentionActionsTests(AttentionAnalyticsTestCase):
    def post(self, action, payload, category=TEAMS, program=None):
        pk = (program or self.program).pk
        return self.client.post(
            f"/programs/{pk}/manager-overview/attention/{category}/{action}/",
            payload,
            format="json",
        )

    def selection(self, ids=None, search="", excluded=None):
        return {
            "search": search,
            "selection": {
                "mode": "all" if ids is None else "selected",
                "ids": ids or [],
                "excluded_ids": excluded or [],
            },
        }

    def preview(self, payload=None, category=TEAMS):
        response = self.post("reminder-preview", payload or self.selection(), category)
        self.assertEqual(response.status_code, 200, response.data)
        return response.data

    def confirm(self, preview, category=TEAMS):
        response = self.post("reminders", {"token": preview["token"]}, category)
        self.assertEqual(response.status_code, 200, response.data)
        return response.data

    def test_access_for_all_actions_is_program_scoped(self):
        member = self.participant().user
        for user in (None, member):
            self.client.force_authenticate(user)
            for action in ("reminder-preview", "reminders", "export"):
                with self.subTest(user=user, action=action):
                    response = self.post(action, self.selection())
                    self.assertIn(response.status_code, (401, 403))
        self.client.force_authenticate(self.manager)
        self.assertEqual(
            self.post("export", self.selection(), program=self.other_program).status_code,
            403,
        )
        self.assertEqual(ProgramAttentionReminder.objects.count(), 0)

    def test_preview_is_read_only_and_all_respects_search_and_exclusions_beyond_page(
        self,
    ):
        ids = [self.participant(first_name="Target").user_id for _ in range(28)]
        self.participant(first_name="Other")
        self.participant(program=self.other_program, first_name="Target")
        result = self.preview(self.selection(search="Target", excluded=[ids[0]]))
        self.assertEqual(result["selected_count"], 27)
        self.assertEqual(result["recipient_count"], 27)
        self.assertEqual(ProgramAttentionReminder.objects.count(), 0)
        self.assertEqual(ProgramAttentionReminderBatch.objects.count(), 0)

    def test_never_visited_uses_both_activity_fields_and_not_email_activation(self):
        unseen = self.participant(is_active=False).user
        visited = self.participant(last_login=self.program.datetime_started).user
        active = self.participant(last_activity=self.program.datetime_started).user
        result = self.preview(self.selection([unseen.pk, visited.pk, active.pk]), VISITS)
        self.assertEqual(result["recipient_count"], 1)
        self.confirm(result, VISITS)
        self.assertEqual(ProgramAttentionReminder.objects.get().recipient_id, unseen.pk)

    def test_confirmation_is_idempotent_and_queues_after_commit(self):
        self.participant()
        preview = self.preview()
        with patch("partner_programs.services.attention_actions.queue_batch") as queue:
            with self.captureOnCommitCallbacks(execute=True):
                result = self.confirm(preview)
                self.assertEqual(queue.call_count, 0)
            self.assertEqual(queue.call_count, 1)
            self.assertEqual(self.confirm(preview), result)
            self.assertEqual(queue.call_count, 1)
        self.assertEqual(result["created_count"], 1)
        self.assertEqual(ProgramAttentionReminder.objects.count(), 1)

    def test_token_tampering_expiry_and_cross_actor_rejected(self):
        self.participant()
        preview = self.preview()
        self.assertEqual(
            self.post("reminders", {"token": preview["token"] + "tamper"}).status_code,
            400,
        )
        with patch(
            "partner_programs.services.attention_actions.signing.loads",
            side_effect=signing.SignatureExpired,
        ):
            self.assertEqual(
                self.post("reminders", {"token": preview["token"]}).status_code, 400
            )
        other = create_user(prefix="other-manager")
        self.program.managers.add(other)
        self.client.force_authenticate(other)
        self.assertEqual(
            self.post("reminders", {"token": preview["token"]}).status_code, 400
        )
        self.assertEqual(ProgramAttentionReminder.objects.count(), 0)

    def test_preview_cannot_be_used_for_another_category_or_program(self):
        self.participant()
        preview = self.preview()
        self.other_program.managers.add(self.manager)
        self.assertEqual(
            self.post("reminders", {"token": preview["token"]}, VISITS).status_code, 400
        )
        self.assertEqual(
            self.post(
                "reminders", {"token": preview["token"]}, program=self.other_program
            ).status_code,
            400,
        )

    def test_participants_resolved_since_preview_are_skipped(self):
        user = self.participant().user
        preview = self.preview(category=VISITS)
        user.last_login = self.program.datetime_started
        user.save(update_fields=["last_login"])
        result = self.confirm(preview, VISITS)
        self.assertEqual(result, {"created_count": 0, "skipped_count": 1})
        self.assertFalse(ProgramAttentionReminder.objects.exists())

    def test_experts_receive_one_message_with_only_own_unfinished_works(self):
        first_work, second_work = (
            self.work(name="Work A"),
            self.work(name="Work B"),
        )
        self.work(name="No expert")
        other = create_rate_expert(program=self.program)
        self.assignment(first_work)
        self.assignment(second_work)
        self.assignment(first_work, expert=other)
        self.score(first_work, expert=other)
        preview = self.preview(category=WORKS)
        self.assertEqual(preview["selected_count"], 3)
        self.assertEqual(preview["recipient_count"], 1)
        self.assertEqual(preview["without_assignments_count"], 1)
        self.confirm(preview, WORKS)
        notice = ProgramAttentionReminder.objects.get()
        self.assertEqual(notice.recipient_id, self.expert.pk)
        self.assertIn("Work A", notice.text)
        self.assertIn("Work B", notice.text)
        self.assertNotIn("No expert", notice.text)

    def test_open_mode_without_real_assignments_has_export_only(self):
        self.open_mode()
        self.work()
        preview = self.preview(category=WORKS)
        self.assertIsNone(preview["token"])
        self.assertEqual(preview["recipient_count"], 0)
        self.assertEqual(self.post("export", self.selection(), WORKS).status_code, 200)

    def test_completed_and_new_assignments_after_preview_are_not_sent(self):
        first_work, second_work = self.work(name="Complete later"), self.work(name="Keep")
        self.assignment(first_work)
        self.assignment(second_work)
        preview = self.preview(category=WORKS)
        self.score(first_work)
        other = create_rate_expert(program=self.program)
        self.assignment(second_work, expert=other)
        self.confirm(preview, WORKS)
        notice = ProgramAttentionReminder.objects.get()
        self.assertEqual(notice.recipient_id, self.expert.pk)
        self.assertIn("Keep", notice.text)
        self.assertNotIn("Complete later", notice.text)

    def test_reassigned_work_does_not_leak_into_message_for_retained_assignment(self):
        first_work, second_work = self.work(name="Removed assignment"), self.work(
            name="Retained assignment"
        )
        first = self.assignment(first_work)
        self.assignment(second_work)
        preview = self.preview(category=WORKS)
        first.delete()
        other = create_rate_expert(program=self.program)
        self.assignment(first_work, expert=other)
        self.confirm(preview, WORKS)
        notice = ProgramAttentionReminder.objects.get()
        self.assertNotIn("Removed assignment", notice.text)
        self.assertIn("Retained assignment", notice.text)

    def test_notification_feed_and_read_are_recipient_owned(self):
        user = self.participant().user
        self.confirm(self.preview())
        notice = ProgramAttentionReminder.objects.get()
        self.client.force_authenticate(user)
        response = self.client.get("/core/notifications/")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["unread_count"], 1)
        self.assertEqual(response.data["results"][0]["kind"], "program_attention")
        url = f"/core/notifications/program_attention:{notice.pk}/read/"
        self.assertEqual(self.client.post(url).status_code, 200)
        self.assertEqual(self.client.post(url).data["unread_count"], 0)
        self.client.force_authenticate(self.manager)
        self.assertEqual(self.client.post(url).status_code, 404)

    def test_private_email_is_escaped_and_retries_are_idempotent(self):
        user = self.participant(first_name="<script>alert(1)</script>").user
        self.confirm(self.preview())
        notice = ProgramAttentionReminder.objects.get()
        deliver_reminder(notice.pk)
        deliver_reminder(notice.pk)
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, [user.email])
        self.assertEqual(mail.outbox[0].cc, [])
        self.assertNotIn("<script>", mail.outbox[0].alternatives[0][0])
        notice.refresh_from_db()
        self.assertEqual(notice.delivery_status, "sent")
        self.assertEqual(notice.attempts, 1)

    def test_email_failure_is_saved_and_periodic_dispatch_recovers(self):
        self.participant()
        self.confirm(self.preview())
        notice = ProgramAttentionReminder.objects.get()
        with patch(
            "partner_programs.services.attention_actions.EmailMultiAlternatives.send",
            side_effect=RuntimeError("smtp"),
        ):
            with self.assertRaises(RuntimeError):
                deliver_reminder(notice.pk)
        notice.refresh_from_db()
        self.assertEqual(notice.attempts, 1)
        self.assertEqual(notice.delivery_status, "queued")
        with patch("partner_programs.tasks.send_attention_reminder.delay") as queue:
            self.assertEqual(dispatch_pending_reminders(), 1)
            queue.assert_called_once_with(notice.pk)
        deliver_reminder(notice.pk)
        notice.refresh_from_db()
        self.assertEqual(notice.delivery_status, "sent")

    def test_email_is_cancelled_if_problem_resolved_before_worker(self):
        user = self.participant().user
        self.confirm(self.preview(category=VISITS), VISITS)
        notice = ProgramAttentionReminder.objects.get()
        user.last_login = self.program.datetime_started
        user.save(update_fields=["last_login"])
        deliver_reminder(notice.pk)
        notice.refresh_from_db()
        self.assertEqual(notice.delivery_status, "cancelled")
        self.assertEqual(len(getattr(mail, "outbox", [])), 0)

    def test_rollback_removes_both_platform_notice_and_email_outbox(self):
        self.participant()
        preview = self.preview()
        with patch("partner_programs.services.attention_actions.queue_batch") as queue:
            with self.captureOnCommitCallbacks(execute=True):
                with transaction.atomic():
                    self.confirm(preview)
                    transaction.set_rollback(True)
            queue.assert_not_called()
        self.assertFalse(ProgramAttentionReminder.objects.exists())

    def test_export_matches_search_all_pages_and_keeps_formula_names_literal(self):
        for _ in range(27):
            self.participant(first_name="=HYPERLINK(1)", city="Target")
        self.participant(first_name="Other")
        response = self.post("export", self.selection(search="=HYPERLINK(1)"))
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.content))
        sheet = book.active
        self.assertEqual(sheet.max_row, 28)
        self.assertEqual(sheet.cell(2, 2).data_type, "s")
        self.assertTrue(sheet.cell(2, 2).value.startswith("=HYPERLINK(1)"))

    def test_explicit_selection_excludes_other_program_and_stale_rows(self):
        own = self.participant().user
        foreign = self.participant(program=self.other_program).user
        result = self.preview(self.selection([own.pk, foreign.pk]))
        self.assertEqual(result["recipient_count"], 1)
        result = self.post("export", self.selection([own.pk, foreign.pk]))
        self.assertEqual(load_workbook(io.BytesIO(result.content)).active.max_row, 2)

    def test_selection_and_unknown_category_validation(self):
        for selection in (
            {"mode": "selected", "ids": []},
            {"mode": "all", "ids": [1]},
            {"mode": "selected", "ids": [-1]},
            {"mode": "other"},
        ):
            self.assertEqual(
                self.post("reminder-preview", {"selection": selection}).status_code, 400
            )
        self.assertEqual(
            self.post("export", self.selection(), "not-first_work-category").status_code,
            404,
        )
        self.assertFalse(ProgramAttentionReminder.objects.exists())

    def test_missing_email_is_excluded_and_reported(self):
        user = self.participant().user
        type(user).objects.filter(pk=user.pk).update(email="")
        preview = self.preview()
        self.assertEqual(preview["unavailable_count"], 1)
        self.assertEqual(preview["recipient_count"], 0)
        self.assertIsNone(preview["token"])

    def test_broker_outage_keeps_committed_notification_for_periodic_recovery(self):
        self.participant()
        preview = self.preview()
        with patch(
            "partner_programs.tasks.dispatch_attention_reminders.delay",
            side_effect=RuntimeError("broker"),
        ):
            with self.captureOnCommitCallbacks(execute=True):
                self.confirm(preview)
        self.assertEqual(ProgramAttentionReminder.objects.get().delivery_status, "queued")

    def test_selection_limit_is_validated_before_creating_a_preview(self):
        self.participant()
        with patch("partner_programs.services.attention_actions.MAX_SELECTION", 0):
            response = self.post("reminder-preview", self.selection())
        self.assertEqual(response.status_code, 400)
        self.assertFalse(ProgramAttentionReminderBatch.objects.exists())


@skipUnless(connection.vendor == "postgresql", "PostgreSQL row locks")
class AttentionReminderConcurrencyTests(TransactionTestCase):
    def test_simultaneous_confirmation_creates_one_batch_and_one_notice(self):
        manager = create_user(prefix="concurrent-attention-manager")
        program = create_partner_program()
        program.managers.add(manager)
        user = create_user(prefix="concurrent-attention-member")
        create_program_member(program, user=user)
        preview = preview_reminder(
            program,
            TEAMS,
            manager,
            {"search": "", "selection": {"mode": "all", "ids": [], "excluded_ids": []}},
        )

        def confirm():
            close_old_connections()
            try:
                return confirm_reminder(program, TEAMS, manager, preview["token"])
            finally:
                close_old_connections()

        with patch("partner_programs.services.attention_actions.queue_batch"):
            with ThreadPoolExecutor(max_workers=2) as pool:
                results = list(pool.map(lambda _: confirm(), range(2)))
        self.assertEqual(results, [{"created_count": 1, "skipped_count": 0}] * 2)
        self.assertEqual(ProgramAttentionReminderBatch.objects.count(), 1)
        self.assertEqual(ProgramAttentionReminder.objects.count(), 1)
