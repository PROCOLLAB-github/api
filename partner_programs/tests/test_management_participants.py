from datetime import datetime
from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from partner_programs.tests.helpers import (
    create_partner_program,
    create_program_member,
    create_user,
)


class ManagementParticipantsTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.manager = create_user()
        cls.program = create_partner_program()
        cls.program.managers.add(cls.manager)
        cls.rows = []
        for name, region, activity in [
            ("Анна", " Москва ", "2026-10-18T21:00:00+00:00"),
            ("Анна", "Москва", "2026-10-19T20:59:59+00:00"),
            ("Анна", "Москва", "2026-10-19T21:00:00+00:00"),
            ("Анна", "Казань", "2026-10-19T12:00:00+00:00"),
            ("Иван", "  ", None),
            ("Иван", None, None),
        ]:
            user = create_user(first_name=name, region=region)
            user.__class__.objects.filter(pk=user.pk).update(
                last_activity=datetime.fromisoformat(activity) if activity else None
            )
            cls.rows.append(create_program_member(cls.program, user=user))
        create_program_member(
            create_partner_program(), user=create_user(region="Чужой регион")
        )
        cls.root = f"/programs/{cls.program.pk}/management/participants/"

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.manager)

    def test_intersecting_filters_use_full_inclusive_moscow_days(self):
        response = self.client.get(
            self.root,
            {
                "search": "Анна",
                "region_scope": "selected",
                "region": "Москва",
                "activity": "period",
                "activity_from": "2026-10-19",
                "activity_to": "2026-10-19",
                "limit": 1,
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 2)
        self.assertEqual(response.data["total_count"], 6)
        self.assertEqual(response.data["results"][0]["region"], "Москва")
        self.assertEqual(response.data["region_options"], ["Казань", "Москва"])
        second = self.client.get(
            self.root, {"activity": "period", "activity_to": "2026-10-18"}
        )
        self.assertEqual(second.data["count"], 0)

    def test_missing_activity_and_region_are_explicit_filters(self):
        response = self.client.get(
            self.root, {"region_scope": "without_region", "activity": "never"}
        )
        self.assertEqual(response.data["count"], 2)
        self.assertTrue(
            all(
                row["region"] is None and row["last_activity"] is None
                for row in response.data["results"]
            )
        )

    def test_invalid_filters_fail_for_list_and_export(self):
        for query in [
            {"activity": "period"},
            {"activity": "period", "activity_from": "broken"},
            {
                "activity": "period",
                "activity_from": "2026-10-20",
                "activity_to": "2026-10-19",
            },
            {"activity": "never", "activity_from": "2026-10-19"},
            {"region_scope": "selected"},
            {"region_scope": "wrong"},
            {"region": "Москва"},
            {"activity": "active", "activity_to": "2026-10-19"},
            {"activity": "wrong"},
            {"order": "user__email"},
        ]:
            for suffix in ("", "export/"):
                with self.subTest(query=query, suffix=suffix):
                    self.assertEqual(
                        self.client.get(self.root + suffix, query).status_code, 400
                    )

    def test_activity_presence_filters_match_selection_and_export(self):
        selection = self.root + "selection/"
        for mode, expected in [
            ("active", self.rows[:4]),
            ("inactive", self.rows[4:]),
            ("never", self.rows[4:]),
        ]:
            with self.subTest(mode=mode):
                query = {"activity": mode, "limit": 1, "offset": 1}
                response = self.client.get(self.root, query)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.data["count"], len(expected))
                selected = self.client.get(selection, query)
                self.assertEqual(selected.status_code, 200)
                self.assertEqual(selected.data["ids"], [p.pk for p in expected])
                sheet = load_workbook(
                    BytesIO(self.client.get(self.root + "export/", query).content)
                ).active
                self.assertEqual(sheet.max_row, len(expected) + 1)
                for row in range(2, sheet.max_row + 1):
                    self.assertEqual(
                        isinstance(sheet.cell(row, 6).value, datetime), mode == "active"
                    )

    def test_activity_sort_is_global_stable_and_nulls_always_last(self):
        # Одинаковые даты сортируются по id, отсутствие активности идет в конце.
        user = self.rows[3].user
        user.last_activity = datetime.fromisoformat("2026-10-19T20:59:59+00:00")
        user.save(update_fields=["last_activity"])
        for order, indices in [
            ("activity_asc", [0, 1, 3, 2, 4, 5]),
            ("activity_desc", [2, 1, 3, 0, 4, 5]),
        ]:
            with self.subTest(order=order):
                expected = [self.rows[i].pk for i in indices]
                full = self.client.get(self.root, {"order": order})
                self.assertEqual([p["id"] for p in full.data["results"]], expected)
                page = self.client.get(
                    self.root, {"order": order, "limit": 2, "offset": 2}
                )
                self.assertEqual([p["id"] for p in page.data["results"]], expected[2:4])
                filtered = self.client.get(
                    self.root,
                    {"order": order, "region_scope": "selected", "region": "Москва"},
                )
                self.assertEqual(
                    [p["id"] for p in filtered.data["results"]],
                    [pk for pk in expected if pk in [p.pk for p in self.rows[:3]]],
                )
                sheet = load_workbook(
                    BytesIO(
                        self.client.get(self.root + "export/", {"order": order}).content
                    )
                ).active
                self.assertEqual(
                    [sheet.cell(i, 3).value for i in range(2, 8)],
                    [self.rows[i].user.email for i in indices],
                )

    def test_export_uses_all_filtered_rows_ignoring_pagination_and_keeps_typed_dates(
        self,
    ):
        response = self.client.get(
            self.root + "export/",
            {"region_scope": "selected", "region": "Москва", "limit": 1, "offset": 1},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("no-store", response["Cache-Control"])
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual([sheet.cell(row, 1).value for row in range(2, 5)], [1, 2, 3])
        self.assertEqual(sheet.cell(2, 6).value, datetime(2026, 10, 19))
        self.assertEqual(sheet.cell(2, 6).data_type, "d")

    def test_region_matches_profile_picker_city_with_legacy_region_fallback(self):
        user = self.rows[0].user
        user.city = " Татарстан "
        user.region = "Устаревшее значение"
        user.save(update_fields=["city", "region"])
        response = self.client.get(
            self.root, {"region_scope": "selected", "region": "Татарстан"}
        )
        self.assertEqual(response.data["count"], 1)
        self.assertEqual(response.data["results"][0]["region"], "Татарстан")
        self.assertNotIn("Устаревшее значение", response.data["region_options"])

    def test_export_neutralizes_user_formulas(self):
        user = self.rows[0].user
        user.first_name = '=HYPERLINK("bad")'
        user.save(update_fields=["first_name"])
        sheet = load_workbook(
            BytesIO(self.client.get(self.root + "export/").content)
        ).active
        self.assertTrue(sheet.cell(2, 2).value.startswith("="))
        self.assertEqual(sheet.cell(2, 2).data_type, "s")

    def test_export_permission_and_empty_result(self):
        empty = self.client.get(self.root + "export/", {"search": "Не существует"})
        self.assertEqual(load_workbook(BytesIO(empty.content)).active.max_row, 1)
        for user, status in [(create_user(), 403), (None, 401)]:
            self.client.force_authenticate(user)
            for suffix in ("", "export/"):
                self.assertEqual(self.client.get(self.root + suffix).status_code, status)
