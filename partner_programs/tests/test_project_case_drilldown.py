"""Контракт, права, XLSX и число запросов нового режима аналитики кейсов."""

import io
from urllib.parse import parse_qs, unquote, urlsplit

from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from openpyxl import load_workbook
from rest_framework.test import APIClient

from partner_programs.models import (
    PartnerProgramField,
    PartnerProgramFieldValue,
    PartnerProgramProject,
)
from partner_programs.serializers.project_case_drilldown import ProjectCaseRowSerializer
from partner_programs.services.exports import (
    BASE_COLUMNS,
    build_program_projects_export_file,
)
from partner_programs.services.case_analytics import build_case_analytics
from partner_programs.services.project_case_drilldown import case_project_rows
from partner_programs.tests.helpers import (
    create_partner_program,
    create_program_member,
    create_program_project,
    create_project,
    create_user,
)
from partner_programs.tests.test_case_fields import create_case_field
from project_rates.tests.helpers import create_rate_expert
from projects.models import Collaborator
from projects.serializers import ProjectListSerializer


class ProjectCaseDrilldownTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.program = create_partner_program(is_competitive=True)
        cls.manager = create_user(password=None)
        cls.program.managers.add(cls.manager)
        cls.field = create_case_field(
            cls.program,
            options=["A", "B", "Empty", "Без выбранного кейса", "Кейс + &=100%"],
        )

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.manager)

    def url(self, export=False, program=None):
        program = program or self.program
        suffix = "export-projects" if export else "projects"
        return f"/programs/{program.pk}/{suffix}/"

    def get(self, scope="selected", *, name="A", export=False, program=None, **params):
        query = {"view": "case_analytics", "case_scope": scope, **params}
        if scope == "selected":
            query["case_name"] = name
        return self.client.get(self.url(export, program), query)

    def link(
        self, value=None, *, program=None, field=None, project=None, submitted=False
    ):
        link = create_program_project(
            program or self.program,
            project=project or create_project(leader=self.manager),
            submitted=submitted,
        )
        if value is not None:
            # Исторические невалидные значения читаются без ослабления правил записи.
            PartnerProgramFieldValue.objects.bulk_create(
                [
                    PartnerProgramFieldValue(
                        program_project=link, field=field or self.field, value_text=value
                    )
                ]
            )
        return link

    def page(self, *args, **kwargs):
        response = self.get(*args, **kwargs)
        self.assertEqual(response.status_code, 200, response.data)
        return response.json()

    def workbook(self, *args, **kwargs):
        response = self.get(*args, export=True, **kwargs)
        self.assertEqual(response.status_code, 200, response.content[:300])
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        book = load_workbook(io.BytesIO(response.content), data_only=False)
        self.assertEqual(book.sheetnames, ["Проекты"])
        self.addCleanup(book.close)
        rows = list(book.active.iter_rows(values_only=True))
        self.assertEqual(
            list(rows[0]),
            [
                "№",
                "Название проекта",
                "Кейс",
                "Ссылка на презентацию",
                "Лидер",
                "Регион",
                "Размер команды",
                "Сдача решения",
                "Дата привязки к программе",
                "Дата сдачи решения",
            ],
        )
        return book, rows[1:], response

    def test_selected_projection_and_full_bucket_metrics(self):
        create_program_member(self.program, user=self.manager)
        member = create_user(password=None)
        create_program_member(self.program, user=member)
        unregistered = create_user(password=None)
        project = create_project(
            leader=self.manager,
            name="Selected",
            region="Москва",
            presentation_address="https://example.org/project.pdf",
            draft=True,
        )
        Collaborator.objects.create(project=project, user=member)
        Collaborator.objects.create(project=project, user=unregistered)
        link = self.link("A", project=project, submitted=True)
        self.link("A")
        self.link("B")
        page = self.page(search=" Selected ")
        self.assertEqual(page["count"], 1)
        self.assertEqual(page["selection"], {"scope": "selected", "case_name": "A"})
        self.assertTrue(page["cases_configured"])
        self.assertTrue(page["submission_applicable"])
        self.assertEqual(
            page["case_metrics"],
            {
                "projects_total": 2,
                "participants_total": 2,
                "submitted": 1,
                "not_submitted": 1,
            },
        )
        row = page["results"][0]
        self.assertEqual(
            set(row),
            {
                "program_project_id",
                "project",
                "case",
                "leader",
                "team_size",
                "linked_at",
                "submitted",
                "submitted_at",
            },
        )
        self.assertEqual(row["program_project_id"], link.pk)
        self.assertEqual(
            row["project"],
            {
                "id": project.pk,
                "name": "Selected",
                "region": "Москва",
                "presentation_address": "https://example.org/project.pdf",
            },
        )
        self.assertEqual(row["case"], {"kind": "selected", "name": "A"})
        self.assertEqual(
            row["leader"],
            {
                "user_id": self.manager.pk,
                "full_name": "Program User",
            },
        )
        # Создание Project также добавляет лидера в Collaborator. Сохраняем
        # формулу существующего экспорта (1 + строки), включая эту запись.
        self.assertEqual(row["team_size"], 4)
        self.assertTrue(row["submitted"])
        self.assertEqual(parse_datetime(row["submitted_at"]), link.datetime_submitted)
        self.assertEqual(parse_datetime(row["linked_at"]), link.datetime_created)

    def test_classifier_parity_for_missing_blank_obsolete_and_nonexact_values(self):
        values = [
            "A",
            "B",
            "Без выбранного кейса",
            None,
            "",
            " \t\n",
            "old",
            "a",
            " A",
            "A ",
        ]
        links = [self.link(value) for value in values]
        expected_without = {link.pk for link in links[3:]}
        overview = build_case_analytics(self.program)
        without = self.page("without_case")
        self.assertEqual(without["case_metrics"], overview["without_case"])
        self.assertEqual(
            {r["program_project_id"] for r in without["results"]}, expected_without
        )
        for row in without["results"]:
            self.assertEqual(row["case"], {"kind": "without_case", "name": None})
        for bucket in overview["items"]:
            page = self.page(name=bucket["name"])
            self.assertEqual(page["count"], bucket["projects_total"])
            self.assertEqual(
                page["case_metrics"], {k: v for k, v in bucket.items() if k != "name"}
            )
        all_page = self.page("all")
        self.assertEqual(all_page["count"], len(values))
        self.assertIsNone(all_page["case_metrics"])
        self.assertEqual(all_page["selection"], {"scope": "all", "case_name": None})
        _, rows, _ = self.workbook("without_case")
        self.assertEqual({r[1] for r in rows}, {link.project.name for link in links[3:]})
        self.assertTrue(all(r[2] == "Без выбранного кейса" for r in rows))
        _, literal_rows, _ = self.workbook(name="Без выбранного кейса")
        self.assertEqual([r[1] for r in literal_rows], [links[2].project.name])
        self.assertEqual(
            PartnerProgramFieldValue.objects.get(program_project=links[-1]).value_text,
            "A ",
        )

    def test_without_case_without_definition_or_options_and_empty_program(self):
        program = create_partner_program()
        program.managers.add(self.manager)
        page = self.page("without_case", program=program)
        self.assertEqual(page["count"], 0)
        self.assertFalse(page["cases_configured"])
        self.link(program=program)
        page = self.page("without_case", program=program)
        self.assertEqual(page["count"], 1)
        self.assertFalse(page["cases_configured"])
        _, rows, _ = self.workbook("without_case", program=program)
        self.assertEqual(len(rows), 1)
        field = create_case_field(program)
        PartnerProgramField.objects.filter(pk=field.pk).update(options="")
        page = self.page("without_case", program=program)
        self.assertTrue(page["cases_configured"])
        self.assertEqual(page["case_metrics"]["projects_total"], 1)
        self.assertEqual(self.get(program=program).status_code, 400)
        _, rows, _ = self.workbook("without_case", program=program)
        self.assertEqual(len(rows), 1)

    def test_removed_option_is_rejected_and_its_links_move_to_without_case(self):
        link = self.link("A")
        self.assertEqual(self.page()["count"], 1)
        PartnerProgramField.objects.filter(pk=self.field.pk).update(options="B")
        for export in (False, True):
            self.assertEqual(self.get(export=export).status_code, 400)
        page = self.page("without_case")
        self.assertEqual(
            [row["program_project_id"] for row in page["results"]], [link.pk]
        )
        self.assertEqual(page["case_metrics"]["projects_total"], 1)
        _, rows, _ = self.workbook("without_case")
        self.assertEqual(rows[0][2], "Без выбранного кейса")
        self.assertEqual(
            PartnerProgramFieldValue.objects.get(program_project=link).value_text, "A"
        )

    def test_same_project_in_two_programs_has_independent_case_submission_and_dates(self):
        other = create_partner_program(is_competitive=True)
        other.managers.add(self.manager)
        other_field = create_case_field(other, options=["A", "B"])
        shared = create_project(
            leader=self.manager, presentation_address="https://example.org/shared"
        )
        link_a = self.link("A", project=shared, submitted=True)
        link_b = self.link("B", project=shared, program=other, field=other_field)
        created = timezone.now() - timezone.timedelta(days=10)
        PartnerProgramProject.objects.filter(pk=link_b.pk).update(
            datetime_created=created
        )
        row_a = self.page()["results"][0]
        row_b = self.page(name="B", program=other)["results"][0]
        self.assertNotEqual(row_a["program_project_id"], row_b["program_project_id"])
        self.assertEqual(row_a["program_project_id"], link_a.pk)
        self.assertEqual(row_b["program_project_id"], link_b.pk)
        self.assertEqual(row_a["project"], row_b["project"])
        self.assertTrue(row_a["submitted"])
        self.assertFalse(row_b["submitted"])
        self.assertIsNone(row_b["submitted_at"])
        self.assertEqual(parse_datetime(row_b["linked_at"]), created)
        self.assertEqual(self.page(program=other)["count"], 0)
        _, rows_a, _ = self.workbook()
        _, rows_b, _ = self.workbook(name="B", program=other)
        self.assertEqual(rows_a[0][2], "A")
        self.assertEqual(rows_b[0][2], "B")
        self.assertEqual(rows_a[0][7], "Сдано")
        self.assertEqual(rows_b[0][7], "Не сдано")
        self.assertIsNone(rows_b[0][9])
        # Даже исторически некорректные значения другой программы игнорируются.
        self.link("A", field=other_field)
        self.assertEqual(self.page("without_case")["count"], 1)

    def test_pagination_search_exact_encoded_option_and_full_exports(self):
        name = "Кейс + &=100%"
        links = [
            self.link(
                name,
                project=create_project(
                    leader=self.manager, name=f"{'Find' if i < 3 else 'Project'} {i:02d}"
                ),
            )
            for i in range(30)
        ]
        self.link("B")
        self.link()
        PartnerProgramProject.objects.filter(pk__in=[x.pk for x in links]).update(
            datetime_created=timezone.now()
        )
        first = self.page(name=name)
        self.assertEqual(len(first["results"]), 25)
        self.assertEqual(first["count"], 30)
        self.assertIsNone(first["previous"])
        self.assertEqual(parse_qs(urlsplit(first["next"]).query)["case_name"], [name])
        second = self.client.get(first["next"]).json()
        ids = [r["program_project_id"] for r in first["results"] + second["results"]]
        self.assertEqual(ids, [link.pk for link in links])
        self.assertIsNone(second["next"])
        self.assertIsNotNone(second["previous"])
        search = self.page(name=name, search=" fInD ", limit=2)
        self.assertEqual(search["count"], 3)
        self.assertEqual(search["case_metrics"]["projects_total"], 30)
        self.assertEqual(self.page(name=name, search="missing")["count"], 0)
        self.assertEqual(self.page(name=name, offset=100)["results"], [])
        _, selected, _ = self.workbook(name=name)
        self.assertEqual([r[1] for r in selected], [link.project.name for link in links])
        _, all_rows, _ = self.workbook("all")
        self.assertEqual(len(all_rows), 32)
        self.assertEqual(len({r[1] for r in all_rows}), 32)
        self.assertEqual([r[0] for r in all_rows], list(range(1, 33)))

    def test_noncompetitive_null_presentation_and_historical_submission_dates(self):
        self.program.is_competitive = False
        self.program.save(update_fields=["is_competitive"])
        historical = self.link("A", submitted=True)
        unsent = self.link("A")
        PartnerProgramProject.objects.filter(pk=historical.pk).update(
            datetime_submitted=None
        )
        PartnerProgramProject.objects.filter(pk=unsent.pk).update(
            datetime_submitted=timezone.now()
        )
        page = self.page()
        self.assertFalse(page["submission_applicable"])
        self.assertEqual([r["submitted"] for r in page["results"]], [True, False])
        for row in page["results"]:
            self.assertIsNone(row["submitted_at"])
            self.assertIsNone(row["project"]["presentation_address"])
        _, rows, _ = self.workbook()
        for row in rows:
            self.assertEqual(row[7], "Не требуется")
            self.assertIsNone(row[3])
            self.assertIsNone(row[9])
            self.assertIsNotNone(parse_datetime(row[8]).tzinfo)

    def test_empty_case_returns_zero_metadata_and_header_only_workbook(self):
        self.assertEqual(self.page(name="Empty")["count"], 0)
        self.assertEqual(self.page(name="Empty")["case_metrics"]["projects_total"], 0)
        _, rows, _ = self.workbook(name="Empty")
        self.assertEqual(rows, [])

    def test_query_validation_rejects_invalid_and_ambiguous_scope_without_fallback(self):
        self.link("A")
        invalid = [
            {},
            {"case_scope": "unknown"},
            {"case_scope": "selected"},
            {"case_scope": "selected", "case_name": "a"},
            {"case_scope": "selected", "case_name": " A "},
            {"case_scope": "selected", "case_name": "old"},
            {"case_scope": "selected", "case_name": ""},
            {"case_scope": "selected", "case_name": " \t "},
            {"case_scope": "all", "case_name": "A"},
            {"case_scope": "without_case", "case_name": "A"},
            {"case_scope": "all", "unknown": "1"},
            {"case_scope": "all", "only_submitted": "1"},
        ]
        for export in (False, True):
            for query in invalid:
                with self.subTest(export=export, query=query):
                    response = self.client.get(
                        self.url(export), {"view": "case_analytics", **query}
                    )
                    self.assertEqual(response.status_code, 400)
                    self.assertNotIn("Content-Disposition", response)
            for query in (
                "view=wrong&case_scope=all",
                "case_scope=selected&case_name=A",
                "view=case_analytics&view=wrong&case_scope=all",
                "view=case_analytics&case_scope=all&case_scope=selected",
                "view=case_analytics&case_scope=selected&case_name=A&case_name=B",
            ):
                self.assertEqual(
                    self.client.get(self.url(export) + "?" + query).status_code, 400
                )
        for params in (
            {"limit": 0},
            {"limit": 101},
            {"limit": "x"},
            {"offset": -1},
            {"offset": "x"},
            {"limit": ""},
            {"limit": "1.5"},
            {"offset": ""},
        ):
            self.assertEqual(self.get(**params).status_code, 400)
        for key in ("search", "limit", "offset", "only_submitted"):
            for value in ("", "1"):
                self.assertEqual(self.get(export=True, **{key: value}).status_code, 400)
        self.assertEqual(self.page(limit=100)["count"], 1)

    def test_permission_matrix_missing_program_and_read_only_methods(self):
        self.link(
            "A", project=create_project(leader=self.manager, name="Private case project")
        )
        foreign = create_partner_program()
        foreign_manager = create_user(password=None)
        foreign.managers.add(foreign_manager)
        participant = create_program_member(self.program).user
        expert = create_rate_expert(program=self.program)
        outsider = create_user(password=None)
        leader = create_program_member(self.program).user
        teammate = create_program_member(self.program).user
        owned = self.link("A", project=create_project(leader=leader))
        Collaborator.objects.create(project=owned.project, user=teammate)
        staff = create_user(password=None, is_staff=True)
        superuser = create_user(password=None, is_superuser=True)
        for user, expected in (
            (None, 401),
            (participant, 403),
            (expert, 403),
            (outsider, 403),
            (leader, 403),
            (teammate, 403),
            (foreign_manager, 403),
            (self.manager, 200),
            (staff, 200),
            (superuser, 200),
        ):
            self.client.force_authenticate(user)
            for export in (False, True):
                with self.subTest(user=user, export=export):
                    response = self.get(export=export)
                    self.assertEqual(response.status_code, expected)
                    if expected != 200:
                        self.assertNotIn(b"Private case project", response.content)
                        self.assertNotIn("Content-Disposition", response)
                    missing_url = f"/programs/{foreign.pk + 100000}/{'export-projects' if export else 'projects'}/"
                    missing = self.client.get(
                        missing_url, {"view": "case_analytics", "case_scope": "all"}
                    )
                    self.assertEqual(
                        missing.status_code,
                        401
                        if user is None
                        else 404
                        if user in (staff, superuser)
                        else 403,
                    )
        self.client.force_authenticate(self.manager)
        for export in (False, True):
            for method in ("post", "put", "patch", "delete"):
                self.assertEqual(
                    getattr(self.client, method)(self.url(export)).status_code, 405
                )

    def test_legacy_list_and_all_submitted_xlsx_contracts_stay_unchanged(self):
        links = [self.link("A", submitted=i == 0) for i in range(12)]
        legacy = self.client.get(self.url())
        self.assertEqual(legacy.status_code, 200)
        self.assertEqual(set(legacy.data), {"count", "next", "previous", "results"})
        self.assertEqual(len(legacy.data["results"]), 10)
        expected = ProjectListSerializer(
            [x.project for x in reversed(links[-10:])],
            many=True,
            context={"request": legacy.wsgi_request},
        ).data
        self.assertEqual(legacy.data["results"], expected)
        for submitted in (False, True):
            response = self.client.get(
                self.url(True), {"only_submitted": "1"} if submitted else {}
            )
            self.assertEqual(response.status_code, 200)
            actual = load_workbook(io.BytesIO(response.content))
            reference = build_program_projects_export_file(
                program=self.program, only_submitted=submitted
            )
            baseline = load_workbook(io.BytesIO(reference.binary_data))
            self.addCleanup(actual.close)
            self.addCleanup(baseline.close)
            actual_rows = list(actual.active.values)
            self.assertEqual(actual_rows, list(baseline.active.values))
            self.assertEqual(
                list(actual_rows[0]),
                [title for _, title in BASE_COLUMNS] + [self.field.label],
            )
            self.assertEqual(len(actual_rows) - 1, 1 if submitted else 12)

    def test_xlsx_user_values_are_literal_text_and_filename_is_safe(self):
        option = '=HYPERLINK("https://example.org", "case")'
        PartnerProgramField.objects.filter(pk=self.field.pk).update(options=option)
        self.program.name = 'Program / unsafe\r\n"name'
        self.program.save(update_fields=["name"])
        self.manager.first_name = "=1+1"
        self.manager.last_name = ""
        self.manager.save(update_fields=["first_name", "last_name"])
        names = [
            "=1+1",
            "+SUM(1,2)",
            "-1+2",
            "@SUM(1,2)",
            "#N/A",
            "Текст & + %",
            "\t=1+1",
        ]
        url = "https://example.org/presentation?x=1&token=a+b%20c#slide-15"
        for name in names:
            self.link(
                option,
                project=create_project(
                    leader=self.manager,
                    name=name,
                    region="=1+1",
                    presentation_address=url,
                ),
            )
        book, rows, response = self.workbook(name=option)
        self.assertEqual([r[1] for r in rows], names)
        for cells in list(book.active.iter_rows())[1:]:
            for index in (1, 2, 3, 4, 5):
                self.assertEqual(cells[index].data_type, "s")
                self.assertIsNone(cells[index].hyperlink)
            self.assertEqual(cells[2].value, option)
            self.assertEqual(cells[3].value, url)
            self.assertEqual(cells[4].value, "=1+1")
        disposition = unquote(response["Content-Disposition"])
        for unsafe in ("\r", "\n", "/", "\\"):
            self.assertNotIn(unsafe, disposition)
        self.assertIn("projects_case", disposition)
        self.assertIn(".xlsx", disposition)
        for scope, label in (("all", "all_cases"), ("without_case", "without_case")):
            _, _, response = self.workbook(scope)
            self.assertIn(label, response["Content-Disposition"])

    def test_constant_queries_for_one_and_one_hundred_rows_and_no_writes(self):
        counts = {"list": [], "export": [], "overview": []}
        self.link("A")
        collaborator = create_user(password=None)
        for size in (1, 100):
            if size == 100:
                for index in range(99):
                    project = create_project(leader=self.manager, name=f"Load {index}")
                    Collaborator.objects.create(project=project, user=collaborator)
                    self.link("A", project=project)
            for operation in counts:
                with CaptureQueriesContext(connection) as queries:
                    if operation == "overview":
                        payload = build_case_analytics(self.program)
                        self.assertEqual(payload["items"][0]["projects_total"], size)
                    elif operation == "list":
                        self.assertEqual(len(self.page(limit=100)["results"]), size)
                    else:
                        _, rows, _ = self.workbook()
                        self.assertEqual(len(rows), size)
                counts[operation].append(len(queries))
                for query in queries:
                    self.assertNotRegex(
                        query["sql"], r"(?i)\b(INSERT INTO|UPDATE|DELETE FROM)\b"
                    )
        self.assertEqual(counts["overview"], [4, 4])
        self.assertEqual(counts["list"][0], counts["list"][1], counts)
        self.assertEqual(counts["export"][0], counts["export"][1], counts)
        self.assertLessEqual(counts["list"][1], 10, counts)
        self.assertLessEqual(counts["export"][1], 7, counts)
        rows = list(
            case_project_rows(
                self.program, field=self.field, selection={"case_scope": "all"}
            )
        )
        with self.assertNumQueries(0):
            self.assertEqual(len(ProjectCaseRowSerializer(rows, many=True).data), 100)
