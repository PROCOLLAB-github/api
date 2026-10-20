from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from django.db import connections
from django.test import TransactionTestCase, skipUnlessDBFeature
import copy
import io
from datetime import timedelta
from unittest.mock import patch

from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from partner_programs.models import (
    PartnerProgram,
    PartnerProgramUserProfile,
    ProgramRegistration,
    PendingProgramRegistration,
)
from partner_programs.services.native_registration import confirmation_token
from users.models import CustomUser


TEST_REGISTRATION_SETTINGS = {
    "EMAIL_BACKEND": "django.core.mail.backends.locmem.EmailBackend",
    "EMAIL_USER": "no-reply@example.test",
    "CELERY_TASK_ALWAYS_EAGER": True,
    "CELERY_TASK_EAGER_PROPAGATES": True,
    "CELERY_BROKER_URL": "memory://",
    "CELERY_RESULT_BACKEND": "cache+memory://",
    "PASSWORD_HASHERS": ["django.contrib.auth.hashers.MD5PasswordHasher"],
}


@override_settings(**TEST_REGISTRATION_SETTINGS)
class NativeRegistrationTests(TestCase):
    def setUp(self):
        from django.core.cache import cache

        cache.clear()
        self.client = APIClient()
        self.manager = CustomUser.objects.create_user(
            email="manager@example.test",
            password="manager-password",
            first_name="Иван",
            last_name="Иванов",
            birthday="2000-01-01",
            is_active=True,
        )
        self.participant = CustomUser.objects.create_user(
            email="member@example.test",
            password="member-password",
            first_name="Анна",
            last_name="Иванова",
            birthday="2000-01-01",
            is_active=True,
        )
        now = timezone.now()
        self.program = PartnerProgram.objects.create(
            name="Тестовая программа",
            city="Москва",
            tag="тест",
            draft=False,
            datetime_started=now,
            datetime_finished=now + timedelta(days=30),
            datetime_registration_ends=now + timedelta(days=10),
        )
        self.program.managers.add(self.manager)
        self.field = {
            "id": "f" + "a" * 32,
            "label": "Вуз",
            "type": "select",
            "help_text": "",
            "required": True,
            "source": "own",
            "options": ["Первый вуз", "Второй вуз"],
            "allow_other": False,
        }
        self.schema = {"mode": "native", "fields": [self.field]}
        self.config = ProgramRegistration.objects.create(
            program=self.program,
            draft=self.schema,
            published=self.schema,
            published_version=1,
        )
        self.root = f"/programs/{self.program.pk}"

    def data(self):
        return {"version": 1, "answers": {self.field["id"]: "Первый вуз"}}

    def signup(self):
        return {
            **self.data(),
            "first_name": "Мария",
            "last_name": "Иванова",
            "email": "new@example.test",
            "birthday": "2000-01-01",
            "password": "good-secure-pass-2030",
            "repeat_password": "good-secure-pass-2030",
            "accept_policy": True,
            "accept_age": True,
        }

    def post(self, path, data):
        return self.client.post(self.root + path, data, format="json")

    def test_public_form_hides_drafts(self):
        self.assertEqual(self.client.get(self.root + "/registration/").status_code, 200)
        self.config.published = None
        self.config.save()
        self.assertEqual(self.client.get(self.root + "/registration/").status_code, 400)
        self.program.draft = True
        self.program.save()
        self.assertEqual(self.client.get(self.root + "/registration/").status_code, 404)

    def test_organizer_permissions_and_stale_draft(self):
        self.client.force_authenticate(self.participant)
        self.assertEqual(
            self.client.get(self.root + "/management/registration/").status_code, 403
        )
        self.client.force_authenticate(self.manager)
        data = {"revision": 0, "draft": self.schema}
        self.assertEqual(
            self.client.put(
                self.root + "/management/registration/", data, format="json"
            ).status_code,
            200,
        )
        self.assertEqual(
            self.client.put(
                self.root + "/management/registration/", data, format="json"
            ).status_code,
            409,
        )

    def test_draft_missing_options_can_save_but_not_publish(self):
        self.client.force_authenticate(self.manager)
        schema = copy.deepcopy(self.schema)
        schema["fields"][0]["options"] = []
        self.assertEqual(
            self.client.put(
                self.root + "/management/registration/",
                {"revision": 0, "draft": schema},
                format="json",
            ).status_code,
            200,
        )
        self.assertEqual(
            self.post(
                "/management/registration/publish/", {"revision": 1, "draft": schema}
            ).status_code,
            400,
        )
        self.assertEqual(
            self.client.get(self.root + "/registration/").data["fields"][0]["options"],
            self.field["options"],
        )

    def test_public_submit_needs_authenticated_owner(self):
        self.assertEqual(self.post("/registration/", self.data()).status_code, 403)
        self.assertFalse(PartnerProgramUserProfile.objects.exists())
        self.client.force_authenticate(self.participant)
        self.assertEqual(self.post("/registration/", self.data()).status_code, 200)
        self.assertEqual(PartnerProgramUserProfile.objects.get().user, self.participant)

    def test_existing_email_requires_login_and_never_changes_profile(self):
        data = self.signup()
        data["email"] = self.participant.email.upper()
        self.assertEqual(self.post("/registration/signup/", data).status_code, 409)
        self.participant.refresh_from_db()
        self.assertEqual(self.participant.first_name, "Анна")
        self.assertFalse(PartnerProgramUserProfile.objects.exists())

    def test_signup_pending_confirm_single_use(self):
        with self.captureOnCommitCallbacks(execute=True):
            response = self.post("/registration/signup/", self.signup())
        self.assertEqual(response.status_code, 201, response.data)
        self.assertNotIn("access", response.data)
        self.assertFalse(PartnerProgramUserProfile.objects.exists())
        pending = PendingProgramRegistration.objects.get()
        self.assertFalse(pending.user.is_active)
        self.assertIsNone(pending.user.verification_date)
        self.assertIsNone(pending.user.onboarding_stage)
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn(
            f"/program-registration/{self.program.pk}#confirmation=", mail.outbox[0].body
        )
        result = self.post(
            "/registration/confirm/", {"token": confirmation_token(pending)}
        )
        self.assertEqual(result.status_code, 200, result.data)
        self.assertIn("access", result.data)
        profile = PartnerProgramUserProfile.objects.get()
        self.assertEqual(profile.partner_program_data, self.data()["answers"])
        self.assertEqual(profile.registration_schema, self.schema)
        self.assertEqual(
            self.post(
                "/registration/confirm/", {"token": confirmation_token(pending)}
            ).status_code,
            409,
        )

    def test_dob_policy_password_are_required(self):
        for key, value in (
            ("birthday", None),
            ("accept_policy", False),
            ("repeat_password", "different"),
            ("first_name", "bad123"),
        ):
            data = self.signup()
            data[key] = value
            self.assertEqual(
                self.post("/registration/signup/", data).status_code, 400, key
            )
        self.assertFalse(PendingProgramRegistration.objects.exists())

    def test_invalid_values_and_stale_publication(self):
        self.client.force_authenticate(self.participant)
        for answers in ({}, {self.field["id"]: "Невалидный вуз"}, {"unknown": "value"}):
            self.assertEqual(
                self.post(
                    "/registration/", {"version": 1, "answers": answers}
                ).status_code,
                400,
            )
        self.assertEqual(
            self.post("/registration/", {**self.data(), "version": 2}).status_code, 409
        )
        self.assertFalse(PartnerProgramUserProfile.objects.exists())

    def test_closed_new_enrollment_blocked_existing_idempotent(self):
        self.client.force_authenticate(self.participant)
        self.assertEqual(self.post("/registration/", self.data()).status_code, 200)
        self.program.datetime_registration_ends = timezone.now() - timedelta(seconds=1)
        self.program.save()
        self.assertEqual(self.post("/registration/", self.data()).status_code, 200)
        self.client.force_authenticate(None)
        self.assertEqual(
            self.post("/registration/signup/", self.signup()).status_code, 409
        )
        self.assertEqual(PartnerProgramUserProfile.objects.count(), 1)

    def test_confirmation_after_closure_activates_account_without_enrollment(self):
        self.post("/registration/signup/", self.signup())
        pending = PendingProgramRegistration.objects.get()
        self.program.datetime_registration_ends = timezone.now() - timedelta(seconds=1)
        self.program.save()
        result = self.post(
            "/registration/confirm/", {"token": confirmation_token(pending)}
        )
        self.assertEqual(result.status_code, 200)
        self.assertFalse(result.data["enrolled"])
        self.assertFalse(PartnerProgramUserProfile.objects.exists())
        pending.user.refresh_from_db()
        self.assertTrue(pending.user.is_active)

    def test_resend_receipt_limited_and_invalidates_old_link(self):
        response = self.post("/registration/signup/", self.signup())
        pending = PendingProgramRegistration.objects.get()
        old = confirmation_token(pending)
        self.assertEqual(
            self.post(
                "/registration/resend/", {"token": response.data["receipt"]}
            ).status_code,
            409,
        )
        PendingProgramRegistration.objects.filter(pk=pending.pk).update(
            last_sent_at=timezone.now() - timedelta(minutes=2)
        )
        self.assertEqual(
            self.post(
                "/registration/resend/", {"token": response.data["receipt"]}
            ).status_code,
            200,
        )
        self.assertEqual(
            self.post("/registration/confirm/", {"token": old}).status_code, 409
        )
        self.assertEqual(
            self.post("/registration/resend/", {"token": "bogus"}).status_code, 400
        )

    def test_expired_confirmation_cannot_issue_tokens(self):
        self.post("/registration/signup/", self.signup())
        pending = PendingProgramRegistration.objects.get()
        with patch("django.core.signing.time.time", return_value=1):
            token = confirmation_token(pending)
        self.assertEqual(
            self.post("/registration/confirm/", {"token": token}).status_code, 409
        )
        self.assertFalse(PartnerProgramUserProfile.objects.exists())

    def test_archive_retains_old_answers_and_export(self):
        self.client.force_authenticate(self.participant)
        self.post("/registration/", self.data())
        self.client.force_authenticate(self.manager)
        response = self.post(
            "/management/registration/publish/",
            {"revision": 0, "draft": {"mode": "native", "fields": []}},
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertIn(self.field["id"], response.data["answered_field_ids"])
        profile = PartnerProgramUserProfile.objects.get()
        self.assertEqual(profile.partner_program_data, self.data()["answers"])
        result = self.client.get(self.root + "/management/participants/export/")
        book = load_workbook(io.BytesIO(result.content))
        self.assertEqual(list(book["Анкеты"].values)[1][-2:], ("Вуз", "Первый вуз"))

    def test_import_preview_duplicates_blanks_and_errors_does_not_save(self):
        self.client.force_authenticate(self.manager)
        result = self.client.post(
            self.root + "/management/registration/import/",
            {
                "file": SimpleUploadedFile(
                    "list.csv", "Первый\n\nПервый\nВторой\n".encode()
                )
            },
            format="multipart",
        )
        self.assertEqual(result.status_code, 200, result.data)
        self.assertEqual(
            result.data, {"options": ["Первый", "Второй"], "duplicates": 1, "blanks": 1}
        )
        for filename, content in (
            ("list.xlsx", b"broken"),
            ("list.csv", b"one,two\n"),
            ("list.exe", b"bad"),
        ):
            result = self.client.post(
                self.root + "/management/registration/import/",
                {"file": SimpleUploadedFile(filename, content)},
                format="multipart",
            )
            self.assertEqual(result.status_code, 400)
        self.config.refresh_from_db()
        self.assertEqual(self.config.draft, self.schema)

    def test_xlsx_one_column_and_formulas_blocked(self):
        self.client.force_authenticate(self.manager)
        for value, expected in (("Университет", 200), ("=1+1", 400)):
            book = Workbook()
            book.active.append([value])
            buffer = io.BytesIO()
            book.save(buffer)
            response = self.client.post(
                self.root + "/management/registration/import/",
                {"file": SimpleUploadedFile("list.xlsx", buffer.getvalue())},
                format="multipart",
            )
            self.assertEqual(response.status_code, expected, response.data)

    def test_external_connection_request_cannot_self_approve(self):
        self.client.force_authenticate(self.manager)
        response = self.post(
            "/management/registration/connection/",
            {"url": "https://example.test/registration", "comment": "Форма Тильды"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["external_status"], "requested")
        self.config.refresh_from_db()
        self.assertEqual(
            self.config.connection_website_url, "https://example.test/registration"
        )
        self.assertEqual(self.config.external_url, "")
        self.assertEqual(
            self.post(
                "/management/registration/publish/",
                {"revision": 1, "draft": {"mode": "external", "fields": []}},
            ).status_code,
            400,
        )
        self.assertEqual(
            self.post(
                "/management/registration/connection/",
                {
                    "url": "https://example.test",
                    "comment": "",
                    "external_status": "connected",
                },
            ).status_code,
            400,
        )

    def test_external_request_optional_comment_and_contact_are_saved_for_admin(self):
        self.client.force_authenticate(self.manager)
        response = self.post(
            "/management/registration/connection/", {"url": "https://event.example.test"}
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["connection_comment"], "")
        self.assertEqual(response.data["connection_contact"], "")
        response = self.post(
            "/management/registration/connection/",
            {
                "url": "https://event.example.test",
                "comment": "Нужна регистрация на чемпионат",
                "contact": "@event_organizer",
            },
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.config.refresh_from_db()
        self.assertEqual(self.config.connection_website_url, "https://event.example.test")
        self.assertEqual(self.config.connection_contact, "@event_organizer")
        self.assertEqual(self.config.connection_comment, "Нужна регистрация на чемпионат")
        self.assertEqual(self.config.external_url, "")
        self.assertEqual(response.data["connection_contact"], "@event_organizer")
        self.assertEqual(response.data["connection_website_url"], "https://event.example.test")
        self.assertEqual(self.config.requested_by, self.manager)

    def test_external_request_invalid_details_do_not_write(self):
        self.client.force_authenticate(self.manager)
        for data in (
            {"url": "not-a-website"},
            {"url": "https://example.test", "contact": "x" * 501},
            {"url": "https://example.test", "comment": "x" * 2001},
        ):
            with self.subTest(data=list(data)):
                response = self.post("/management/registration/connection/", data)
                self.assertEqual(response.status_code, 400, response.data)
        self.config.refresh_from_db()
        self.assertEqual(self.config.revision, 0)
        self.assertEqual(self.config.connection_website_url, "")
        self.assertEqual(self.config.connection_contact, "")

    def test_connection_details_migration_preserves_old_request_and_form_url(self):
        import importlib
        from django.apps import apps
        from django.db import connection
        from types import SimpleNamespace

        self.config.external_status = "requested"
        self.config.external_url = "https://previous.example.test/form"
        self.config.save()
        migration = importlib.import_module(
            "partner_programs.migrations.0025_registration_connection_details"
        )
        migration.preserve_requested_websites(
            apps, SimpleNamespace(connection=connection)
        )
        self.config.refresh_from_db()
        self.assertEqual(
            self.config.connection_website_url, "https://previous.example.test/form"
        )
        self.assertEqual(self.config.external_url, "https://previous.example.test/form")
        self.config.connection_website_url = "https://event.example.test"
        self.config.save()
        migration.preserve_requested_websites(
            apps, SimpleNamespace(connection=connection)
        )
        self.config.refresh_from_db()
        self.assertEqual(self.config.connection_website_url, "https://event.example.test")

    def test_external_request_requires_manager_and_does_not_change_connected_form(self):
        data = {"url": "https://event.example.test", "contact": "@event_organizer"}
        self.client.force_authenticate(self.participant)
        self.assertEqual(
            self.post("/management/registration/connection/", data).status_code, 403
        )
        self.config.external_status = "connected"
        self.config.external_url = "https://forms.example.test/register"
        self.config.save()
        self.client.force_authenticate(self.manager)
        response = self.post("/management/registration/connection/", data)
        self.assertEqual(response.status_code, 200, response.data)
        self.config.refresh_from_db()
        self.assertEqual(self.config.external_status, "connected")
        self.assertEqual(self.config.external_url, "https://forms.example.test/register")
        self.assertEqual(self.config.connection_website_url, "")
        self.assertEqual(self.config.connection_contact, "")

    def test_legacy_registration_cannot_bypass_published_native_schema(self):
        self.client.force_authenticate(self.participant)
        self.assertEqual(self.post("/register/", {}).status_code, 400)
        self.client.force_authenticate(None)
        self.assertEqual(self.post("/register_new/", self.signup()).status_code, 400)

    def test_new_required_fields_before_confirmation_require_review(self):
        self.post("/registration/signup/", self.signup())
        pending = PendingProgramRegistration.objects.get()
        newer = copy.deepcopy(self.schema)
        newer["fields"].append({**self.field, "id": "f" + "b" * 32, "label": "Курс"})
        self.config.published = newer
        self.config.published_version = 2
        self.config.save()
        result = self.post(
            "/registration/confirm/", {"token": confirmation_token(pending)}
        )
        self.assertEqual(result.status_code, 200, result.data)
        self.assertTrue(result.data["review_required"])
        self.assertFalse(result.data["enrolled"])
        self.assertEqual(result.data["answers"], self.data()["answers"])
        self.assertFalse(PartnerProgramUserProfile.objects.exists())

    def test_region_filter_uses_program_answer_without_overwriting_profile(self):
        schema = {
            "mode": "native",
            "fields": [{**self.field, "source": "regions", "options": []}],
        }
        self.config.published = schema
        self.config.save()
        self.client.force_authenticate(self.participant)
        result = self.post(
            "/registration/", {"version": 1, "answers": {self.field["id"]: "Москва"}}
        )
        self.assertEqual(result.status_code, 200, result.data)
        self.participant.refresh_from_db()
        self.assertFalse(self.participant.city)
        self.client.force_authenticate(self.manager)
        result = self.client.get(
            self.root + "/management/participants/?region_scope=selected&region=Москва"
        )
        self.assertEqual(result.data["count"], 1)
        self.assertEqual(result.data["results"][0]["region"], "Москва")

    def test_invalid_json_document_is_400(self):
        self.assertEqual(
            self.post("/registration/signup/", [{"bad": True}]).status_code, 400
        )

    def test_stale_browser_token_does_not_hide_public_form(self):
        self.client.credentials(HTTP_AUTHORIZATION="Bearer expired")
        result = self.client.get(self.root + "/registration/")
        self.assertEqual(result.status_code, 200)
        self.assertFalse(result.data["authenticated"])

    def test_confirmation_does_not_work_for_another_program(self):
        self.post("/registration/signup/", self.signup())
        pending = PendingProgramRegistration.objects.get()
        original = self.program.pk
        self.program.pk = None
        self.program.save()
        result = self.client.post(
            f"/programs/{self.program.pk}/registration/confirm/",
            {"token": confirmation_token(pending)},
            format="json",
        )
        self.assertEqual(result.status_code, 404)
        self.assertFalse(
            PartnerProgramUserProfile.objects.filter(partner_program_id=original).exists()
        )

    def test_list_number_date_and_other_validation(self):
        from partner_programs.services.native_registration import validate_answers
        from rest_framework.exceptions import ValidationError

        for kind, value in (
            ("number", True),
            ("date", "2026-99-99"),
            ("multiselect", ["Невалидный"]),
            ("select", "Другое:    "),
        ):
            field = {**self.field, "type": kind, "allow_other": True}
            with self.assertRaises(ValidationError):
                validate_answers({"fields": [field]}, {field["id"]: value})
        self.assertEqual(
            validate_answers(
                {"fields": [{**self.field, "allow_other": True}]},
                {self.field["id"]: "Другое: Третий вуз"},
            )[self.field["id"]],
            "Другое: Третий вуз",
        )

    def test_native_signup_preserves_password_reset_program_context(self):
        self.post("/registration/signup/", self.signup())
        self.client.post(
            "/auth/reset_password/",
            {"email": self.participant.email, "registration_program_id": self.program.pk},
            format="json",
        )
        self.assertIn("/program-password-reset/?", mail.outbox[-1].body)
        self.assertIn(f"registration_program={self.program.pk}", mail.outbox[-1].body)


@skipUnlessDBFeature("has_select_for_update")
@override_settings(**TEST_REGISTRATION_SETTINGS)
class NativeRegistrationConcurrencyTests(TransactionTestCase):
    setUp = NativeRegistrationTests.setUp
    data = NativeRegistrationTests.data
    signup = NativeRegistrationTests.signup

    def parallel(self, calls):
        gate = Barrier(2)

        def run(call):
            connections.close_all()
            try:
                gate.wait(timeout=10)
                return call()
            finally:
                connections.close_all()

        with ThreadPoolExecutor(max_workers=2) as pool:
            return list(pool.map(run, calls))

    def test_double_signup_creates_one_pending_account(self):
        def signup():
            return (
                APIClient()
                .post(self.root + "/registration/signup/", self.signup(), format="json")
                .status_code
            )

        self.assertEqual(sorted(self.parallel([signup, signup])), [201, 409])
        self.assertEqual(PendingProgramRegistration.objects.count(), 1)
        self.assertFalse(PartnerProgramUserProfile.objects.exists())

    def test_double_enrollment_is_idempotent(self):
        def enroll():
            client = APIClient()
            client.force_authenticate(self.participant)
            return client.post(
                self.root + "/registration/", self.data(), format="json"
            ).status_code

        self.assertEqual(self.parallel([enroll, enroll]), [200, 200])
        self.assertEqual(PartnerProgramUserProfile.objects.count(), 1)

    def test_concurrent_publish_has_one_winner(self):
        def publish():
            client = APIClient()
            client.force_authenticate(self.manager)
            return client.post(
                self.root + "/management/registration/publish/",
                {"revision": 0, "draft": self.schema},
                format="json",
            ).status_code

        self.assertEqual(sorted(self.parallel([publish, publish])), [200, 409])
