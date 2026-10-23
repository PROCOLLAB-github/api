from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APIClient

from partner_programs.models import PartnerProgramField, PartnerProgramFieldValue
from project_rates.models import Criteria, ProjectScore
from project_rates.tests.helpers import (
    create_rate_expert,
    create_rate_program,
    create_rate_project,
    create_rate_user,
    link_project_to_program,
)


class ManagementResultsTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.manager = create_rate_user()
        cls.program = create_rate_program(is_competitive=True, required_evaluators=3)
        cls.program.managers.add(cls.manager)
        cls.program.criterias.all().delete()
        cls.criteria = [
            Criteria.objects.create(
                partner_program=cls.program,
                name=name,
                type=kind,
                min_value=minimum,
                max_value=maximum,
            )
            for name, kind, minimum, maximum in [
                ("Аналитика", "int", 0, 10),
                ("Проработка", "int", 0, 20),
                ("Условия", "bool", None, None),
                ("Комментарий", "str", None, None),
            ]
        ]
        cls.experts = [create_rate_expert(program=cls.program) for _ in range(4)]
        cls.field = PartnerProgramField.objects.create(
            partner_program=cls.program,
            name="case",
            label="Кейс",
            field_type="select",
            is_required=True,
            show_filter=True,
            options='Сбер|Т-Банк|Кейс "А"',
        )
        cls.projects = []
        for name, case in [
            ("Город", "Сбер"),
            ("Маршрут", "Сбер"),
            ("Кампус", "Сбер"),
            ("Река", "Сбер"),
            ("Район", "Сбер"),
            ("Без оценок", "Сбер"),
            ("Ноль", "Т-Банк"),
            ("Черновик", "Сбер"),
        ]:
            project = create_rate_project()
            project.name = name
            project.save()
            cls.program.users.add(
                project.leader, through_defaults={"partner_program_data": {}}
            )
            link = link_project_to_program(cls.program, project)
            PartnerProgramFieldValue.objects.create(
                program_project=link, field=cls.field, value_text=case
            )
            type(link).objects.filter(pk=link.pk).update(submitted=name != "Черновик")
            cls.projects.append(project)
        for index, pairs in enumerate(
            [
                [(9, 18), (7, 14), (8, 16)],
                [(8, 14), (7, 15), (8, 15)],
                [(8, 14), (7, 15), (8, 15)],
                [(7, 13), (7, 14), (7, 14)],
                [(9, 18), (8, 18)],
                [],
                [(0, 0), (0, 0), (0, 0)],
                [],
            ]
        ):
            for expert, pair in zip(cls.experts, pairs):
                cls.score(cls.projects[index], expert, pair)

    @classmethod
    def score(cls, project, expert, pair):
        for criterion, value in zip(cls.criteria, [*pair, "False", ""]):
            ProjectScore.objects.create(
                criteria=criterion, user=expert, project=project, value=str(value)
            )

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.manager)

    def url(self, export=False, pk=None):
        return reverse(
            "partner_programs:management-results" + ("-export" if export else ""),
            kwargs={"pk": pk or self.program.pk},
        )

    def rows(self, **query):
        response = self.client.get(self.url(), query)
        self.assertEqual(response.status_code, 200, response.data)
        return [p for g in response.data["groups"] for p in g["projects"]]

    def test_evaluation_detail_uses_same_canonical_scores_as_results(self):
        from partner_programs.models import PartnerProgramProject

        snapshot = self.client.get(self.url()).data
        for expected in self.rows():
            response = self.client.get(
                f'/programs/{self.program.pk}/management/projects/{expected["id"]}/evaluation/'
            )
            self.assertEqual(response.status_code, 200, response.data)
            actual = response.data["result"]
            self.assertIsNone(actual["place"])
            for key in (
                "project_id",
                "name",
                "completed",
                "required",
                "ready",
                "evaluation_complete",
                "score",
                "averages",
                "experts",
            ):
                self.assertEqual(actual[key], expected[key], key)
            self.assertEqual(response.data["criteria"], snapshot["criteria"])
            self.assertIn("can_rate", response.data)
            self.assertIn("experts", response.data)
        link = PartnerProgramProject.objects.get(
            project=self.projects[-1], partner_program=self.program
        )
        response = self.client.get(
            f"/programs/{self.program.pk}/management/projects/{link.pk}/evaluation/"
        )
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.data["submitted"])
        self.assertFalse(response.data["result"]["ready"])
        self.assertIsNone(response.data["result"]["place"])

    def test_rounding_ties_case_ranks_and_pending(self):
        rows = {p["name"]: p for p in self.rows()}
        self.assertEqual((rows["Город"]["score"], rows["Город"]["place"]), (24, 1))
        self.assertEqual((rows["Кампус"]["score"], rows["Кампус"]["place"]), (22.33, 2))
        self.assertEqual(rows["Маршрут"]["place"], 2)
        self.assertEqual(rows["Река"]["place"], 4)
        self.assertEqual((rows["Район"]["score"], rows["Район"]["place"]), (26.5, None))
        self.assertEqual((rows["Ноль"]["score"], rows["Ноль"]["place"]), (0, 1))
        self.assertIsNone(rows["Без оценок"]["score"])
        self.assertNotIn("Черновик", rows)

    def test_sort_and_search_preserve_rank(self):
        rows = self.rows(case_scope="selected", case_name="Сбер", order="asc")
        self.assertEqual([p["place"] for p in rows if p["ready"]], [4, 2, 2, 1])
        self.assertEqual(self.rows(search="РЕКА")[0]["place"], 4)
        self.assertEqual(len(self.rows(case_scope="selected", case_name="Т-Банк")), 1)
        self.assertEqual(self.client.get(self.url(), {"order": "bad"}).status_code, 400)

    def test_partial_evaluation_excluded_and_surplus_retained(self):
        ProjectScore.objects.create(
            criteria=self.criteria[0],
            project=self.projects[0],
            user=self.experts[3],
            value="10",
        )
        row = self.rows(search="Город")[0]
        self.assertEqual((row["score"], row["completed"]), (24, 3))
        ProjectScore.objects.filter(user=self.experts[3]).delete()
        self.score(self.projects[0], self.experts[3], (0, 0))
        row = self.rows(search="Город")[0]
        self.assertEqual((row["score"], row["completed"]), (18, 4))
        type(self.program).objects.filter(pk=self.program.pk).update(
            required_evaluators=5
        )
        row = self.rows(search="Город")[0]
        self.assertEqual((row["score"], row["place"]), (18, None))

    def test_access_and_foreign_program(self):
        for user, expected in [
            (None, 401),
            (self.experts[0], 403),
            (create_rate_user(), 403),
        ]:
            self.client.force_authenticate(user)
            for export in (False, True):
                self.assertEqual(self.client.get(self.url(export)).status_code, expected)
        self.client.force_authenticate(self.manager)
        foreign = create_rate_program()
        self.assertEqual(self.client.get(self.url(pk=foreign.pk)).status_code, 403)

    def test_xlsx_types_formulas_scope_and_text_injection(self):
        self.projects[0].name = "=1+1"
        self.projects[0].save()
        response = self.client.get(self.url(True), {"search": "=1+1"})
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book.sheetnames, ["Рейтинг", "Оценки экспертов"])
        rating, experts = book.worksheets
        self.assertEqual(rating.max_row, 2)
        self.assertEqual(rating["B2"].data_type, "s")
        self.assertEqual(rating["B2"].value, "=1+1")
        self.assertEqual(experts["F2"].data_type, "n")
        self.assertEqual(experts["H2"].value, "=SUM(F2:G2)")
        self.assertIn("ROUND(AVERAGEIF(", rating["H2"].value)
        self.assertIn('"Полная"', rating["H2"].value)
        self.assertEqual(experts["I2"].value, "Нет")
        response = self.client.get(self.url(True), {"search": "nothing", "scope": "all"})
        self.assertEqual(
            load_workbook(BytesIO(response.content)).worksheets[0].max_row, 8
        )

    def test_no_numeric_criteria_produces_no_score(self):
        program = create_rate_program(required_evaluators=1)
        program.managers.add(self.manager)
        program.criterias.all().delete()
        criterion = Criteria.objects.create(
            partner_program=program, name="Условия", type="bool"
        )
        project = create_rate_project()
        program.users.add(project.leader, through_defaults={"partner_program_data": {}})
        link = link_project_to_program(program, project)
        type(link).objects.filter(pk=link.pk).update(submitted=True)
        ProjectScore.objects.create(
            criteria=criterion, project=project, user=self.experts[0], value="True"
        )
        data = self.client.get(self.url(pk=program.pk)).data
        self.assertFalse(data["has_numeric_criteria"])
        row = data["groups"][0]["projects"][0]
        self.assertIsNone(row["score"])
        self.assertIsNone(row["place"])
        self.assertTrue(row["evaluation_complete"])
        self.assertFalse(row["ready"])
        self.assertEqual(row["completed"], 1)
        filtered = self.client.get(self.url(pk=program.pk), {"status": "complete"}).data
        self.assertEqual(len(filtered["groups"][0]["projects"]), 1)
        self.assertEqual(
            self.client.get(self.url(pk=program.pk), {"status": "pending"}).data[
                "groups"
            ],
            [],
        )
        response = self.client.get(self.url(True, pk=program.pk))
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book["Рейтинг"]["G2"].value, "Оценивание завершено")
        self.assertIsNone(book["Рейтинг"]["A2"].value)
        self.assertIsNone(book["Рейтинг"]["F2"].value)
        type(program).objects.filter(pk=program.pk).update(required_evaluators=2)
        row = self.client.get(self.url(pk=program.pk)).data["groups"][0]["projects"][0]
        self.assertFalse(row["evaluation_complete"])
        self.assertIsNone(row["score"])
        response = self.client.get(self.url(True, pk=program.pk))
        self.assertEqual(
            load_workbook(BytesIO(response.content))["Рейтинг"]["G2"].value,
            "Ожидает оценок",
        )

    def test_status_filter_and_export_keep_original_ranks_and_all_scope(self):
        complete = self.rows(status="complete")
        pending = self.rows(status="pending")
        self.assertEqual(len(complete), 5)
        self.assertEqual({p["name"] for p in pending}, {"Район", "Без оценок"})
        self.assertTrue(all(p["evaluation_complete"] for p in complete))
        self.assertFalse(any(p["evaluation_complete"] for p in pending))
        self.assertEqual(self.rows(status="complete", search="Река")[0]["place"], 4)
        for export in (False, True):
            self.assertEqual(
                self.client.get(self.url(export), {"status": "bad"}).status_code, 400
            )
        response = self.client.get(self.url(True), {"status": "pending"})
        self.assertEqual(load_workbook(BytesIO(response.content))["Рейтинг"].max_row, 3)
        response = self.client.get(
            self.url(True), {"status": "pending", "search": "nothing", "scope": "all"}
        )
        self.assertEqual(load_workbook(BytesIO(response.content))["Рейтинг"].max_row, 8)

    def test_comment_only_no_criteria_and_invalid_values_do_not_fake_completion(self):
        self.program = create_rate_program(required_evaluators=3)
        self.program.managers.add(self.manager)
        self.program.criterias.all().delete()
        project = create_rate_project()
        self.program.users.add(
            project.leader, through_defaults={"partner_program_data": {}}
        )
        link = link_project_to_program(self.program, project)
        type(link).objects.filter(pk=link.pk).update(submitted=True)
        row = self.rows()[0]
        self.assertFalse(row["evaluation_complete"])
        criterion = Criteria.objects.create(
            partner_program=self.program, name="Комментарий", type="str"
        )
        for expert in self.experts[:3]:
            ProjectScore.objects.create(
                criteria=criterion, project=project, user=expert, value=""
            )
        row = self.rows()[0]
        self.assertTrue(row["evaluation_complete"])
        self.assertIsNone(row["score"])
        self.assertIsNone(row["place"])
        self.assertFalse(row["ready"])
        self.program = create_rate_program(required_evaluators=1)
        self.program.managers.add(self.manager)
        self.program.criterias.all().delete()
        criterion = Criteria.objects.create(
            partner_program=self.program, name="Условия", type="bool"
        )
        project = create_rate_project()
        self.program.users.add(
            project.leader, through_defaults={"partner_program_data": {}}
        )
        link = link_project_to_program(self.program, project)
        type(link).objects.filter(pk=link.pk).update(submitted=True)
        historical = ProjectScore.objects.create(
            criteria=criterion, project=project, user=self.experts[0], value="False"
        )
        ProjectScore.objects.filter(pk=historical.pk).update(value="")
        self.assertFalse(self.rows()[0]["evaluation_complete"])

    def test_invalid_historical_value_does_not_break_export(self):
        ProjectScore.objects.filter(
            criteria=self.criteria[0], project=self.projects[0], user=self.experts[0]
        ).update(value="bad")
        self.assertEqual(self.rows(search="Город")[0]["completed"], 2)
        self.assertEqual(self.client.get(self.url(True)).status_code, 200)
