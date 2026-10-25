import copy
from io import BytesIO
from urllib.parse import parse_qs, urlsplit

from django.db import connection
from django.test import TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook
from rest_framework.test import APIClient

from files.models import UserFile
from partner_programs import solution_forms
from partner_programs.solution_models import ProgramSolution, ProgramSolutionAnswer
from partner_programs.solution_schema import form_data
from partner_programs.tests.helpers import create_user
from partner_programs.tests.test_solution_forms import call, new_field
from project_rates.tests.test_solution_evaluation import EvaluationFixture


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class SolutionFieldFilterTests(EvaluationFixture, TestCase):
    def setUp(self):
        self.seed_evaluation()
        self.root = f"/programs/{self.program.pk}/management/solution-analytics/"

    def page(self, **query):
        response = call(self.manager, "get", self.root + "solutions/", query)
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response["Cache-Control"], "private, no-store")
        self.assert_schema("ManagerSolutions", response.data)
        return response.data

    def ids(self, **query):
        return [
            row["application"]["solution_id"] for row in self.page(**query)["results"]
        ]

    def book(self, **query):
        response = call(
            self.manager, "get", self.root + "export/", {"kind": "solutions", **query}
        )
        self.assertEqual(response.status_code, 200, getattr(response, "data", None))
        self.assertEqual(response["Cache-Control"], "private, no-store")
        return load_workbook(BytesIO(response.content), data_only=False)

    def answer(self, solution, kind, value):
        if value is None:
            # API удаляет ответ при null; в БД value NOT NULL сохраняется.
            ProgramSolutionAnswer.objects.filter(
                solution=solution, field_id=self.fields[kind]["id"]
            ).delete()
            return
        ProgramSolutionAnswer.objects.update_or_create(
            solution=solution,
            field_id=self.fields[kind]["id"],
            defaults={"value": value},
        )

    def test_no_field_filter_and_explicit_all_preserve_existing_selection(self):
        draft = self.another_solution(submitted=False)
        expected = [self.solution.pk, draft.pk]
        self.assertEqual(self.ids(), expected)
        self.assertEqual(self.ids(answer_state="all"), expected)
        self.assertEqual(
            self.ids(field_id=self.fields["number"]["id"], answer_state="all"), expected
        )
        self.assertEqual(
            [row[0].value for row in list(self.book()["Решения"].rows)[1:]], expected
        )

    def test_zero_and_false_are_present_and_absent_answer_is_missing(self):
        draft = self.another_solution(submitted=False)
        for kind in ("number", "checkbox"):
            query = dict(field_id=self.fields[kind]["id"])
            with self.subTest(kind=kind):
                self.assertEqual(
                    self.ids(**query, answer_state="present"), [self.solution.pk]
                )
                self.assertEqual(self.ids(**query, answer_state="missing"), [draft.pk])
                self.assertEqual(self.ids(**query, answer_state="not_applicable"), [])

    def test_null_empty_and_whitespace_follow_existing_empty_value(self):
        draft = self.another_solution(submitted=False)
        for value in (None, "", " \t\n", "Текст"):
            self.answer(draft, "textarea", value)
            with self.subTest(value=value):
                query = dict(field_id=self.fields["textarea"]["id"])
                self.assertEqual(
                    draft.pk in self.ids(**query, answer_state="present"),
                    value == "Текст",
                )
                self.assertEqual(
                    draft.pk in self.ids(**query, answer_state="missing"),
                    value != "Текст",
                )

    def test_new_required_field_is_not_applicable_to_old_submission(self):
        draft = self.another_solution(submitted=False)
        before = copy.deepcopy(self.solution.submission_snapshot)
        solution_forms.save_form(
            program_id=self.program.pk,
            actor=self.manager,
            **self.payload(
                fields=[
                    *form_data(self.program)["fields"],
                    new_field("number", required=True, position=50),
                ]
            ),
        )
        added = form_data(self.program)["fields"][-1]["id"]
        self.assertEqual(
            self.ids(field_id=added, answer_state="not_applicable"), [self.solution.pk]
        )
        self.assertEqual(self.ids(field_id=added, answer_state="missing"), [draft.pk])
        ProgramSolutionAnswer.objects.create(solution=draft, field_id=added, value=0)
        self.assertEqual(self.ids(field_id=added, answer_state="present"), [draft.pk])
        self.solution.refresh_from_db()
        self.assertEqual(self.solution.submission_snapshot, before)

    def test_historical_snapshot_field_is_selectable_without_current_definition(self):
        # Имитируем историческую запись из миграции, не разрешение редактировать сдачу.
        snapshot = copy.deepcopy(self.solution.submission_snapshot)
        historical = {
            **snapshot["form"]["fields"][1],
            "id": 9000000,
            "label": "Архивное поле",
        }
        snapshot["form"]["fields"].append(historical)
        snapshot["answers"].append(
            dict(field_id=historical["id"], value="Архивный ответ")
        )
        ProgramSolution.objects.filter(pk=self.solution.pk).update(
            submission_snapshot=snapshot
        )
        draft = self.another_solution(submitted=False)
        self.assertEqual(
            self.ids(field_id=historical["id"], answer_state="present"),
            [self.solution.pk],
        )
        self.assertEqual(
            self.ids(field_id=historical["id"], answer_state="not_applicable"), [draft.pk]
        )
        book = self.book(field_id=historical["id"], answer_state="present")
        self.assertIn(
            historical["id"], [row[1].value for row in list(book["Поля формы"].rows)[1:]]
        )

    def test_submitted_answers_and_search_use_snapshot_over_late_live_values(self):
        ProgramSolutionAnswer.objects.filter(
            solution=self.solution, field_id=self.fields["number"]["id"]
        ).update(value="")
        ProgramSolution.objects.filter(pk=self.solution.pk).update(
            name="Позднее live имя"
        )
        name = self.solution.submission_snapshot["name"]
        self.assertEqual(
            self.ids(
                search=name, field_id=self.fields["number"]["id"], answer_state="present"
            ),
            [self.solution.pk],
        )
        self.assertEqual(self.ids(search="Позднее live имя"), [])

    def test_field_filter_intersects_case_submission_attention_and_search(self):
        draft = self.another_solution(submitted=False)
        self.answer(draft, "number", 0)
        query = dict(field_id=self.fields["number"]["id"], answer_state="present")
        self.assertEqual(
            self.ids(
                **query,
                case_scope="selected",
                case_name="Задача 1",
                scope="awaiting_evaluation",
                submission="submitted",
            ),
            [self.solution.pk],
        )
        self.assertEqual(
            self.ids(
                **query, case_scope="without_case", scope="not_submitted", search="Вторая"
            ),
            [draft.pk],
        )
        self.assertEqual(
            self.ids(**query, submission="submitted", scope="not_submitted"), []
        )

    def test_pagination_counts_and_links_preserve_filter_and_export_is_unpaginated(self):
        for _ in range(24):
            self.answer(self.another_solution(submitted=False), "number", 0)
        query = dict(field_id=self.fields["number"]["id"], answer_state="present")
        page = self.page(**query, limit=5, offset=20)
        self.assertEqual(page["count"], 25)
        self.assertEqual(len(page["results"]), 5)
        previous = parse_qs(urlsplit(page["previous"]).query)
        self.assertEqual(previous["field_id"], [str(query["field_id"])])
        self.assertEqual(previous["answer_state"], ["present"])
        self.assertEqual(previous["offset"], ["15"])
        self.assertIsNone(page["next"])
        self.assertEqual(self.book(**query)["Решения"].max_row, 26)

    def test_filtered_workbook_scopes_every_related_sheet_and_keeps_zero_false_types(
        self,
    ):
        self.another_solution(submitted=False)
        query = dict(field_id=self.fields["number"]["id"], answer_state="present")
        book = self.book(**query)
        self.assertEqual(
            set(book.sheetnames), {"Решения", "Команды", "Поля формы", "Ответы решения"}
        )
        for sheet in book:
            self.assertEqual(
                {row[0].value for row in list(sheet.rows)[1:]}, {self.solution.pk}
            )
        answers = {row[1].value: row[4] for row in list(book["Ответы решения"].rows)[1:]}
        self.assertEqual(answers[self.fields["number"]["id"]].value, 0)
        self.assertEqual(answers[self.fields["number"]["id"]].data_type, "n")
        self.assertIs(answers[self.fields["checkbox"]["id"]].value, False)
        self.assertEqual(answers[self.fields["checkbox"]["id"]].data_type, "b")

    def test_submitted_export_intersects_filter_and_empty_workbook_remains_valid(self):
        draft = self.another_solution(submitted=False)
        self.answer(draft, "number", 0)
        query = dict(field_id=self.fields["number"]["id"], answer_state="present")
        book = self.book(kind="submitted", **query)
        self.assertEqual(book["Решения"].max_row, 2)
        self.assertEqual(book["Решения"]["A2"].value, self.solution.pk)
        for query in (
            dict(search="Нет такой работы"),
            dict(kind="submitted", submission="not_submitted"),
        ):
            with self.subTest(query=query):
                self.assertTrue(all(sheet.max_row == 1 for sheet in self.book(**query)))

    def test_invalid_field_queries_are_uniform_c00_errors(self):
        foreign = self.another_program_field()
        for query in (
            dict(answer_state="present"),
            dict(answer_state="missing"),
            dict(answer_state="not_applicable"),
            dict(field_id=0),
            dict(field_id=-1),
            dict(field_id=1.5),
            dict(field_id=9007199254740992),
            dict(field_id="NaN"),
            dict(field_id=foreign),
            dict(field_id=9999999),
            dict(field_id=self.fields["number"]["id"], answer_state="answered"),
            dict(field_id=""),
            dict(field_id="true"),
            dict(answer_state=""),
            dict(project_id=1),
        ):
            for suffix in ("solutions/", "export/"):
                with self.subTest(query=query, suffix=suffix):
                    response = call(
                        self.manager,
                        "get",
                        self.root + suffix,
                        {
                            **({"kind": "solutions"} if suffix == "export/" else {}),
                            **query,
                        },
                    )
                    self.assertEqual(response.status_code, 400, response.data)
                    self.assertEqual(response.data["code"], "validation_error")
                    self.assert_schema("Error", response.data)

    def another_program_field(self):
        from partner_programs.tests.test_solution_foundation import solution_program

        return form_data(solution_program())["fields"][0]["id"]

    def test_global_exports_reject_filters_and_accept_previous_default_calls(self):
        for kind in ("program", "scores"):
            self.book(kind=kind)
            for query in (
                dict(search="Работа"),
                dict(field_id=self.fields["number"]["id"]),
                dict(scope="not_submitted"),
                dict(case_scope="without_case"),
            ):
                with self.subTest(kind=kind, query=query):
                    response = call(
                        self.manager,
                        "get",
                        self.root + "export/",
                        dict(kind=kind, **query),
                    )
                    self.assertEqual(response.status_code, 400, response.data)
        for query in (
            dict(limit=20),
            dict(offset=0),
            dict(case_scope="selected", case_name="Чужой кейс"),
        ):
            response = call(
                self.manager,
                "get",
                self.root + "export/",
                dict(kind="solutions", **query),
            )
            self.assertEqual(response.status_code, 400, response.data)

    def test_filtered_endpoints_keep_real_permissions_before_field_disclosure(self):
        query = dict(field_id=self.fields["number"]["id"], answer_state="present")
        for suffix in ("solutions/", "export/"):
            for actor in (self.captain, self.expert, create_user()):
                response = call(actor, "get", self.root + suffix, query)
                self.assertEqual(response.status_code, 403, response.data)
            self.assertEqual(APIClient().get(self.root + suffix, query).status_code, 401)
            staff = call(
                create_user(is_staff=True),
                "get",
                self.root + suffix,
                {**query, **({"kind": "solutions"} if suffix == "export/" else {})},
            )
            self.assertEqual(staff.status_code, 200)

    def test_noncompetitive_repository_field_needs_no_submission(self):
        self.program.is_competitive = False
        self.program.save(update_fields=["is_competitive"])
        draft = self.another_solution(submitted=False)
        self.answer(draft, "url", "https://example.com/repository")
        query = dict(field_id=self.fields["url"]["id"], answer_state="present")
        page = self.page(**query)
        self.assertFalse(page["submission_applicable"])
        self.assertEqual(page["results"][0]["evaluation_status"], "not_applicable")
        self.assertEqual(self.ids(**query), [draft.pk])
        self.assertEqual(self.book(**query)["Решения"]["H2"].value, "Не требуется")
        draft.refresh_from_db()
        self.assertEqual(draft.status, "draft")

    def test_filtered_reads_are_readonly_with_constant_query_count(self):
        query = dict(field_id=self.fields["number"]["id"], answer_state="present")
        before = list(ProgramSolution.objects.values())
        answers = list(ProgramSolutionAnswer.objects.values())
        files = UserFile.objects.count()
        with CaptureQueriesContext(connection) as list_small:
            self.page(**query)
        with CaptureQueriesContext(connection) as book_small:
            self.book(**query)
        self.assertEqual(list(ProgramSolution.objects.values()), before)
        self.assertEqual(list(ProgramSolutionAnswer.objects.values()), answers)
        self.assertEqual(UserFile.objects.count(), files)
        for _ in range(8):
            self.answer(self.another_solution(submitted=False), "number", 0)
        with CaptureQueriesContext(connection) as list_large:
            self.assertEqual(self.page(**query)["count"], 9)
        with CaptureQueriesContext(connection) as book_large:
            self.assertEqual(self.book(**query)["Решения"].max_row, 10)
        self.assertEqual(len(list_small), len(list_large))
        self.assertEqual(len(book_small), len(book_large))
