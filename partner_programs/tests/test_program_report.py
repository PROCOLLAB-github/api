"""The consolidated report preserves actual data and program boundaries."""

from datetime import datetime, timezone
from io import BytesIO

from django.db import connection
from django.test import TestCase, TransactionTestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook
from rest_framework.test import APIClient

from partner_programs.models import PartnerProgramFieldValue, PartnerProgramUserProfile
from partner_programs.services.program_report import build_program_report
from partner_programs.tests.helpers import (
    create_partner_program,
    create_program_field,
    create_program_member,
    create_program_project,
    create_project,
    create_user,
)
from partner_programs.tests.test_case_fields import create_case_field
from project_rates.models import Criteria, ProjectExpertAssignment, ProjectScore
from project_rates.tests.helpers import create_rate_expert
from projects.models import Collaborator


def rows(book, sheet):
    return list(book[sheet].values)[1:]


class ProgramReportTests(TestCase):
    def setUp(self):
        self.program = create_partner_program(
            is_competitive=True, is_distributed_evaluation=True, max_project_rates=1
        )

    def report(self):
        return load_workbook(BytesIO(build_program_report(self.program)))

    def test_empty_report_includes_all_sheets_and_unscored_criteria(self):
        book = self.report()
        self.assertEqual(
            book.sheetnames,
            [
                "Сводка",
                "Участники",
                "Анкеты",
                "Проекты",
                "Команды",
                "Кейсы",
                "Эксперты",
                "Назначения",
                "Критерии",
                "Оценки",
            ],
        )
        self.assertEqual(rows(book, "Проекты"), [])
        self.assertEqual(rows(book, "Оценки"), [])
        self.assertEqual(len(rows(book, "Критерии")), self.program.criterias.count())
        self.assertEqual(book["Проекты"].freeze_panes, "A2")

    def test_linked_projects_cases_teams_and_program_specific_fields(self):
        case = create_case_field(self.program, options=["Кейс А", "Пустой кейс"])
        extra = create_program_field(self.program, label="Сайт проекта")
        leader = create_user(first_name="Алёна")
        member = create_user()
        link = create_program_project(
            self.program,
            project=create_project(
                leader=leader,
                name="Ёлка",
                presentation_address="https://example.com/slides",
            ),
            submitted=True,
        )
        # Historical duplicates are valid input to the read-only report.
        Collaborator.objects.bulk_create(
            [Collaborator(project=link.project, user=member)]
        )
        PartnerProgramFieldValue.objects.bulk_create(
            [
                PartnerProgramFieldValue(
                    program_project=link, field=case, value_text="Кейс А"
                ),
                PartnerProgramFieldValue(
                    program_project=link, field=extra, value_text="=HYPERLINK(1)"
                ),
            ]
        )
        other = create_partner_program()
        create_program_project(other)
        book = self.report()
        project = rows(book, "Проекты")[0]
        self.assertEqual(len(rows(book, "Проекты")), 1)
        self.assertEqual(project[:3], (link.project_id, link.pk, "Ёлка"))
        self.assertEqual(project[4], "Кейс А")
        self.assertEqual(project[8:11], (2, "https://example.com/slides", "Сдано"))
        self.assertEqual(project[-1], "=HYPERLINK(1)")
        self.assertEqual(book["Проекты"].cell(2, len(project)).data_type, "s")
        self.assertEqual(
            {row[2] for row in rows(book, "Команды")}, {leader.pk, member.pk}
        )
        self.assertEqual([r[4] for r in rows(book, "Команды")].count("Лидер"), 1)
        self.assertIn(("Пустой кейс", 0, 0, 0), rows(book, "Кейсы"))

    def test_participants_deduplicated_but_all_registration_answers_preserved(self):
        user = create_user(first_name="Алёна", last_name="Ёлкина")
        first = create_program_member(
            self.program, user=user, data={"why": "=SUM(1)", "legacy": ["А", "Б"]}
        )
        self.program.data_schema = {"fields": [{"id": "why", "label": "Мотивация"}]}
        self.program.save(update_fields=["data_schema"])
        PartnerProgramUserProfile.objects.filter(pk=first.pk).update(
            datetime_created=datetime(2026, 10, 1, 12, tzinfo=timezone.utc),
        )
        second = create_program_member(self.program, data={"extra": "Другой ответ"})
        create_program_member(create_partner_program(), data={"secret": "Чужая анкета"})
        book = self.report()
        people = rows(book, "Участники")
        self.assertEqual(len(people), 2)
        self.assertIn("Алёна", people[0][1])
        self.assertEqual(people[0][4], datetime(2026, 10, 1, 15))
        answers = rows(book, "Анкеты")
        self.assertEqual({r[0] for r in answers}, {first.pk, second.pk})
        self.assertIn("Мотивация", [r[4] for r in answers])
        self.assertIn('["А", "Б"]', [r[5] for r in answers])
        self.assertNotIn("Чужая анкета", str(answers))
        self.assertEqual(book["Анкеты"]["F2"].data_type, "s")

    def test_scores_zero_comments_assignments_and_cross_program_isolation(self):
        Criteria.objects.filter(partner_program=self.program).delete()
        numeric = Criteria.objects.create(
            partner_program=self.program,
            name="Баллы",
            type="int",
            min_value=0,
            max_value=10,
        )
        comment = Criteria.objects.create(
            partner_program=self.program, name="Комментарий", type="str"
        )
        expert = create_rate_expert(program=self.program)
        link = create_program_project(self.program, submitted=True)
        assignment = ProjectExpertAssignment.objects.create(
            partner_program=self.program, project=link.project, expert=expert.expert
        )
        zero = ProjectScore.objects.create(
            project=link.project, user=expert, criteria=numeric, value="0"
        )
        ProjectScore.objects.create(
            project=link.project, user=expert, criteria=comment, value="=не формула"
        )
        other = create_partner_program()
        foreign_criterion = Criteria.objects.filter(partner_program=other).first()
        foreign_link = create_program_project(other)
        ProjectScore.objects.create(
            project=link.project,
            user=expert,
            criteria=foreign_criterion,
            value="Чужой критерий",
        )
        ProjectScore.objects.create(
            project=foreign_link.project, user=expert, criteria=numeric, value="9"
        )
        book = self.report()
        scores = rows(book, "Оценки")
        self.assertEqual(len(scores), 2)
        self.assertEqual(scores[0][0], zero.pk)
        self.assertEqual(scores[0][8:10], (0, "0"))
        self.assertEqual(book["Оценки"]["I2"].data_type, "n")
        self.assertEqual(book["Оценки"]["I3"].data_type, "s")
        assigned = rows(book, "Назначения")[0]
        self.assertEqual(assigned[:2], (assignment.pk, link.project_id))
        self.assertEqual(assigned[6], "Выполнено")
        self.assertEqual(rows(book, "Эксперты")[0][4:], (1, 1))

    def test_open_evaluation_preserves_real_scores_without_assignments(self):
        self.program.is_distributed_evaluation = False
        self.program.save(update_fields=["is_distributed_evaluation"])
        link = create_program_project(self.program, submitted=True)
        criterion = Criteria.objects.filter(partner_program=self.program).first()
        ProjectScore.objects.create(
            project=link.project,
            user=create_rate_expert(program=self.program),
            criteria=criterion,
            value="Открытая оценка",
        )
        book = self.report()
        self.assertEqual(len(rows(book, "Оценки")), 1)
        self.assertEqual(rows(book, "Назначения"), [])

    def test_exports_beyond_page_size_without_per_row_queries(self):
        link = create_program_project(self.program)
        with CaptureQueriesContext(connection) as small:
            self.report()
        for _ in range(27):
            create_program_member(self.program)
            create_program_project(
                self.program, project=create_project(leader=link.project.leader)
            )
        with CaptureQueriesContext(connection) as large:
            book = self.report()
        self.assertEqual(len(rows(book, "Проекты")), 28)
        self.assertEqual(len(rows(book, "Участники")), 27)
        self.assertLessEqual(len(large), len(small) + 5)


class ProgramReportAccessTests(TransactionTestCase):
    def setUp(self):
        self.client = APIClient()
        self.manager = create_user()
        self.program = create_partner_program(name="=Сохраняем название")
        self.program.managers.add(self.manager)
        self.url = f"/programs/{self.program.pk}/project-analytics/export/"

    def test_auth_and_manager_permission_checked_before_report(self):
        self.assertEqual(self.client.get(self.url).status_code, 401)
        self.client.force_authenticate(create_user())
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.client.force_authenticate(self.manager)
        self.assertEqual(
            self.client.get("/programs/99999999/project-analytics/export/").status_code,
            404,
        )

    def test_read_only_download_has_private_headers_and_literal_program_name(self):
        self.client.force_authenticate(self.manager)
        with CaptureQueriesContext(connection) as captured:
            response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Cache-Control"], "private, no-store")
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertIn("spreadsheetml", response["Content-Type"])
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book["Сводка"]["B3"].value, "=Сохраняем название")
        self.assertEqual(book["Сводка"]["B3"].data_type, "s")
        writes = [
            q["sql"]
            for q in captured
            if q["sql"].lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE"))
        ]
        self.assertEqual(writes, [])
