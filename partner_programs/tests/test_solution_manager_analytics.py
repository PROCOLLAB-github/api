import copy
from io import BytesIO
from unittest.mock import patch

from django.db import connection
from django.test import TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from partner_programs import solution_forms
from partner_programs.models import PartnerProgramProject, PartnerProgramUserProfile
from partner_programs.services.solution_analytics import SolutionAnalyticsContext
from partner_programs.solution_models import (
    ProgramSolution,
    ProgramSolutionMember,
    SolutionInvitation,
)
from partner_programs.solution_schema import form_data
from partner_programs.tests.helpers import (
    create_partner_program,
    create_program_member,
    create_project,
    create_user,
)
from partner_programs.tests.test_solution_forms import call, new_field
from project_rates.models import Criteria
from project_rates.solution_models import SolutionExpertAssignment, SolutionScore
from project_rates.tests.test_solution_evaluation import EvaluationFixture


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class SolutionManagerAnalyticsTests(EvaluationFixture, TestCase):
    def setUp(self):
        self.seed_evaluation()
        self.root = f"/programs/{self.program.pk}/management/solution-analytics/"

    def get(self, suffix="", schema="ManagerAnalytics", query=""):
        response = call(self.manager, "get", self.root + suffix + query)
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response["Cache-Control"], "private, no-store")
        self.assert_schema(schema, response.data)
        return response.data

    def book(self, kind="program"):
        response = call(self.manager, "get", self.root + f"export/?kind={kind}")
        self.assertEqual(response.status_code, 200, getattr(response, "data", None))
        self.assertEqual(response["Cache-Control"], "private, no-store")
        self.assertIn("spreadsheetml.sheet", response["Content-Type"])
        return load_workbook(BytesIO(response.content), data_only=False)

    def test_all_operations_require_program_manager_and_explicit_solution_mode(self):
        suffixes = [
            "",
            "assignments/",
            f"assignments/{self.assignment.pk}/",
            "solutions/",
            "export/",
        ]
        for suffix in suffixes:
            for actor in [self.captain, self.expert, create_user()]:
                response = call(actor, "get", self.root + suffix)
                self.assertEqual(response.status_code, 403, (suffix, response.data))
                self.assert_schema("Error", response.data)
            response = APIClient().get(self.root + suffix)
            self.assertEqual(response.status_code, 401)
            self.assert_schema("Error", response.data)
        other = create_partner_program()
        other.managers.add(self.manager)
        for suffix in suffixes:
            response = call(
                self.manager,
                "get",
                f"/programs/{other.pk}/management/solution-analytics/" + suffix,
            )
            self.assertEqual(response.status_code, 409)
            self.assertEqual(response.data["code"], "solution_mode_unavailable")
        staff = create_user(is_staff=True)
        self.assertEqual(call(staff, "get", self.root).status_code, 200)

    def test_summary_real_solutions_registered_accepted_team_and_case_parity(self):
        member, pending = create_user(city="Москва"), create_user()
        for user in [member, pending]:
            create_program_member(self.program, user=user)
        draft = self.another_solution(submitted=False)
        ProgramSolutionMember.objects.create(solution=draft, user=member)
        SolutionInvitation.objects.create(solution=draft, recipient=pending, role="Роль")
        project = create_project(leader=pending, region="Только регион проекта")
        PartnerProgramProject.objects.bulk_create(
            [PartnerProgramProject(partner_program=self.program, project=project)]
        )
        data = self.get()
        self.assertEqual(data["summary"]["solutions"]["total"], 2)
        self.assertEqual(data["summary"]["participants"]["total"], 4)
        self.assertEqual(data["participant_funnel"]["with_team"], 3)
        self.assertEqual(data["participant_funnel"]["captains"], 2)
        self.assertEqual(data["attention"]["participants_without_team"], 1)
        self.assertEqual(
            data["solution_funnel"],
            dict(
                created=2,
                submission_applicable=True,
                not_submitted=1,
                submitted=1,
                evaluated=0,
            ),
        )
        self.assertEqual(
            data["cases"],
            call(
                self.manager,
                "get",
                f"/programs/{self.program.pk}/management/solution-cases/",
            ).data,
        )
        self.assertNotIn("regions", data["summary"])
        self.assertNotIn("Только регион проекта", str(data))

    def test_distributed_complete_values_zero_false_and_real_quota(self):
        self.program.required_evaluators = 2
        self.program.save(update_fields=["required_evaluators"])
        SolutionExpertAssignment.objects.create(
            partner_program=self.program,
            solution=self.solution,
            expert=self.second.expert,
        )
        self.assertEqual(self.rate().status_code, 200)
        data = self.get()
        self.assertEqual(
            data["evaluation_status"]["assignments"],
            dict(total=2, pending=1, evaluated=1),
        )
        self.assertEqual(data["evaluation_status"]["solutions"]["partially_evaluated"], 1)
        self.assertEqual(data["attention"]["solutions_awaiting_evaluation"], 1)
        self.assertEqual(self.rate(actor=self.second).status_code, 200)
        data = self.get()
        self.assertEqual(data["solution_funnel"]["evaluated"], 1)
        self.assertEqual(data["attention"]["solutions_awaiting_evaluation"], 0)
        detail = self.get(f"assignments/{self.assignment.pk}/", "ManagerAssignmentDetail")
        values = {row["type"]: row for row in detail["scores"]}
        self.assertEqual(values["int"]["value"], "0")
        self.assertEqual(values["bool"]["value"], "False")
        self.assertTrue(values["str"]["is_scored"])
        self.assertEqual(detail["status"], "completed")

    def test_unassigned_score_and_invalid_value_do_not_complete_assignment(
        self,
    ):
        SolutionScore.objects.create(
            solution=self.solution, user=self.second, criteria=self.criteria[0], value="0"
        )
        self.assertEqual(self.get()["solution_funnel"]["evaluated"], 0)
        self.assertEqual(self.rate().status_code, 200)
        SolutionScore.objects.filter(
            solution=self.solution, user=self.expert, criteria=self.criteria[0]
        ).update(value="invalid")
        self.assertEqual(self.get()["evaluation_status"]["assignments"]["evaluated"], 0)

    def test_no_criteria_does_not_complete_assignment(self):
        Criteria.objects.filter(partner_program=self.program).delete()
        self.assertEqual(
            self.get()["evaluation_status"]["assignments"],
            dict(total=1, pending=1, evaluated=0),
        )

    def test_open_first_score_and_noncompetitive_inapplicable_counts(self):
        self.program.is_distributed_evaluation = False
        self.program.save(update_fields=["is_distributed_evaluation"])
        self.assertIsNone(self.get()["evaluation_status"]["assignments"])
        SolutionScore.objects.create(
            solution=self.solution, user=self.second, criteria=self.criteria[0], value="0"
        )
        self.assertEqual(self.get()["solution_funnel"]["evaluated"], 1)
        self.program.is_competitive = False
        self.program.save(update_fields=["is_competitive"])
        draft = self.another_solution(submitted=False)
        data = self.get()
        self.assertEqual(data["summary"]["solutions"]["total"], 2)
        for key in ("submitted", "not_submitted", "evaluated"):
            self.assertIsNone(data["solution_funnel"][key])
        self.assertIsNone(data["evaluation_status"]["solutions"])
        self.assertIsNone(data["participant_funnel"]["submitted_captains"])
        self.assertTrue(
            all(row["submitted_solutions"] is None for row in data["activity"])
        )
        page = self.get("solutions/", "ManagerSolutions", "?scope=not_submitted")
        self.assertEqual(page["count"], 0)
        draft.refresh_from_db()
        self.assertEqual(draft.status, "draft")

    def test_field_coverage_uses_each_applied_form_and_preserves_zero_false(self):
        self.another_solution(submitted=False)
        fields = form_data(self.program)["fields"]
        solution_forms.save_form(
            program_id=self.program.pk,
            actor=self.manager,
            **self.payload(
                fields=[
                    *fields,
                    new_field(
                        required=True, label="Новое обязательное поле", position=50
                    ),
                ]
            ),
        )
        data = self.get()
        coverage = {row["label"]: row for row in data["field_coverage"]}
        new = coverage["Новое обязательное поле"]
        self.assertEqual(
            [
                new[k]
                for k in ("applicable", "required_applicable", "required_unanswered")
            ],
            [1, 1, 1],
        )
        by_type = {row["type"]: row for row in data["field_coverage"]}
        self.assertEqual(by_type["number"]["answered"], 1)
        self.assertEqual(by_type["checkbox"]["answered"], 1)
        self.assertEqual(by_type["file"]["required_unanswered"], 0)
        self.assertEqual(by_type["file"]["required_applicable"], 0)

    def test_waiting_delayed_experts_existing_thresholds_and_not_ready(self):
        now = timezone.now()
        ProgramSolution.objects.filter(pk=self.solution.pk).update(
            submitted_at=now - timezone.timedelta(hours=25)
        )
        SolutionExpertAssignment.objects.filter(pk=self.assignment.pk).update(
            datetime_created=now - timezone.timedelta(hours=26)
        )
        self.program.datetime_evaluation_ends = now + timezone.timedelta(days=4)
        self.program.save(update_fields=["datetime_evaluation_ends"])
        with patch(
            "partner_programs.services.solution_analytics.timezone.now", return_value=now
        ):
            assignment = self.get("assignments/", "ManagerAssignments")["results"][0]
            self.assertEqual(assignment["waiting_seconds"], 25 * 3600)
            delayed = self.get()["attention"]["delayed_experts"]["items"][0]
            self.assertEqual(
                (delayed["overdue_24h"], delayed["overdue_48h"], delayed["severity"]),
                (1, 0, "warning"),
            )
        self.program.datetime_evaluation_ends = now - timezone.timedelta(hours=1)
        self.program.save(update_fields=["datetime_evaluation_ends"])
        self.assertEqual(
            self.get()["attention"]["delayed_experts"]["items"][0]["severity"], "critical"
        )
        self.program.datetime_evaluation_ends = now + timezone.timedelta(days=6)
        self.program.save(update_fields=["datetime_evaluation_ends"])
        self.assertEqual(self.get()["attention"]["delayed_experts"]["total"], 0)
        draft = self.another_solution(submitted=False)
        raw = SolutionExpertAssignment(
            partner_program=self.program, solution=draft, expert=self.expert.expert
        )
        SolutionExpertAssignment.objects.bulk_create([raw])
        rows = self.get("assignments/", "ManagerAssignments")["results"]
        pending = next(
            row for row in rows if row["application"]["solution_id"] == draft.pk
        )
        self.assertEqual(pending["status"], "not_ready")
        self.assertIsNone(pending["waiting_since"])

    def test_activity_30_day_window_real_submission_and_zero_days(self):
        old = self.another_solution()
        ProgramSolution.objects.filter(pk=old.pk).update(
            submitted_at=timezone.now() - timezone.timedelta(days=31)
        )
        activity = self.get()["activity"]
        self.assertEqual(len(activity), 30)
        self.assertEqual(sum(row["submitted_solutions"] for row in activity), 1)
        self.assertTrue(any(row["submitted_solutions"] == 0 for row in activity))

    def test_missing_historical_timestamp_never_invents_wait_duration(self):
        context = SolutionAnalyticsContext(self.program)
        # New submitted rows require a timestamp in DB. Check defensive reader
        # behavior in memory without bypassing or dropping that constraint.
        context.by_id[self.solution.pk].submitted_at = None
        assignment = context.assignment(self.assignment)
        self.assertEqual(assignment["status"], "pending")
        self.assertIsNone(assignment["waiting_since"])
        self.assertIsNone(assignment["waiting_seconds"])

    def test_program_workbook_keeps_defensive_legacy_registration_reading(self):
        PartnerProgramUserProfile.objects.filter(partner_program=self.program).update(
            registration_schema=[], partner_program_data={"legacy": "=HYPERLINK(1)"}
        )
        book = self.book()
        self.assertIn("Анкеты", book.sheetnames)
        self.assertEqual(book["Анкеты"]["F2"].value, "=HYPERLINK(1)")
        self.assertEqual(book["Анкеты"]["F2"].data_type, "s")

    def test_solution_list_filters_snapshot_names_pagination_and_reason(self):
        self.another_solution(submitted=False)
        ProgramSolution.objects.filter(pk=self.solution.pk).update(
            name="Измененное live имя"
        )
        page = self.get("solutions/", "ManagerSolutions", "?limit=1")
        self.assertEqual(page["count"], 2)
        self.assertIn("offset=1", page["next"])
        self.assertEqual(
            page["results"][0]["application"]["name"],
            self.solution.submission_snapshot["name"],
        )
        self.assertEqual(page["results"][0]["waiting_reason"], "no_completed_evaluations")
        waiting = self.get("solutions/", "ManagerSolutions", "?scope=awaiting_evaluation")
        self.assertEqual(
            waiting["count"], self.get()["attention"]["solutions_awaiting_evaluation"]
        )
        self.assertEqual(
            self.get("solutions/", "ManagerSolutions", "?scope=not_submitted")["count"], 1
        )
        self.assertEqual(
            self.get(
                "solutions/",
                "ManagerSolutions",
                "?case_scope=selected&case_name=Задача+1",
            )["count"],
            1,
        )
        self.assertEqual(
            self.get("solutions/", "ManagerSolutions", "?case_scope=without_case")[
                "count"
            ],
            1,
        )
        self.assertEqual(
            self.get("solutions/", "ManagerSolutions", "?search=Измененное")["count"], 0
        )
        self.assertEqual(
            self.get("solutions/", "ManagerSolutions", "?offset=99")["results"], []
        )
        for query in (
            "?case_scope=selected",
            "?case_scope=selected&case_name=bad",
            "?limit=0",
            "?offset=-1",
            "?role=manager",
        ):
            self.assertEqual(
                call(self.manager, "get", self.root + "solutions/" + query).status_code,
                400,
            )

    def test_assignment_scope_foreign_ids_and_foreign_solution_never_leak(self):
        for invalid_id in (0, 9007199254740992, 999999999999999999999999):
            response = call(self.manager, "get", self.root + f"assignments/{invalid_id}/")
            self.assertEqual(response.status_code, 404)
            self.assert_schema("Error", response.data)
        other = EvaluationFixture()
        other.seed_evaluation()
        self.assertEqual(
            call(
                self.manager, "get", self.root + f"assignments/{other.assignment.pk}/"
            ).status_code,
            404,
        )
        SolutionExpertAssignment.objects.bulk_create(
            [
                SolutionExpertAssignment(
                    partner_program=self.program,
                    solution=other.solution,
                    expert=self.expert.expert,
                )
            ]
        )
        rows = self.get("assignments/", "ManagerAssignments")["results"]
        self.assertEqual([row["assignment_id"] for row in rows], [self.assignment.pk])
        self.assertNotEqual(rows[0]["expert"]["expert_id"], rows[0]["expert"]["user_id"])
        self.assertEqual(
            self.get("assignments/", "ManagerAssignments", "?scope=completed")["results"],
            [],
        )
        self.assertEqual(
            len(
                self.get("assignments/", "ManagerAssignments", "?scope=pending")[
                    "results"
                ]
            ),
            1,
        )
        self.assertEqual(self.rate().status_code, 200)
        self.assertEqual(
            self.get("assignments/", "ManagerAssignments", "?scope=pending")["results"],
            [],
        )

    def test_xlsx_program_subsets_snapshot_answers_scores_and_formula_safety(self):
        self.assertEqual(self.rate().status_code, 200)
        SolutionScore.objects.create(
            solution=self.solution, user=self.second, criteria=self.criteria[0], value="1"
        )
        draft = self.another_solution(submitted=False)
        snapshot = copy.deepcopy(self.solution.submission_snapshot)
        snapshot["name"] = "=HYPERLINK(1)"
        ProgramSolution.objects.filter(pk=self.solution.pk).update(
            submission_snapshot=snapshot
        )
        book = self.book()
        self.assertIn("Ответы решения", book.sheetnames)
        self.assertIn("Анкеты", book.sheetnames)
        self.assertEqual(dict(book["Сводка"].values)["Решения"], 2)
        rows = list(book["Решения"].values)
        self.assertEqual({row[0] for row in rows[1:]}, {self.solution.pk, draft.pk})
        for sheet in book:
            self.assertNotIn("ID проекта", next(sheet.values))
            for row in sheet:
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f")
        values = [row[-1] for row in list(book["Ответы решения"].values)[1:]]
        self.assertTrue(any(type(value) is int and value == 0 for value in values))
        self.assertTrue(any(type(value) is bool and value is False for value in values))
        self.assertEqual(len(list(book["Оценки"].values)), 5)
        submitted = self.book("submitted")
        self.assertEqual(
            [row[0] for row in list(submitted["Решения"].values)[1:]], [self.solution.pk]
        )
        self.assertEqual(set(self.book("scores").sheetnames), {"Критерии", "Оценки"})
        self.assertNotIn("Оценки", self.book("solutions").sheetnames)
        response = call(self.manager, "get", self.root + "export/?kind=invalid")
        self.assertEqual(response.status_code, 400)
        self.assert_schema("Error", response.data)

    def test_empty_program_noncompetitive_workbook_and_readonly_constant_queries(self):
        for suffix, schema in [
            ("", "ManagerAnalytics"),
            ("assignments/", "ManagerAssignments"),
            ("solutions/", "ManagerSolutions"),
        ]:
            with CaptureQueriesContext(connection) as queries:
                self.get(suffix, schema)
            self.assertLessEqual(len(queries), 30)
            self.assertFalse(
                any(
                    row["sql"].lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE"))
                    for row in queries
                )
            )
        with CaptureQueriesContext(connection) as before:
            self.get()
        for _ in range(8):
            self.another_solution(submitted=False)
        with CaptureQueriesContext(connection) as after:
            self.get()
        self.assertEqual(len(before), len(after))
        self.program.is_competitive = False
        self.program.save(update_fields=["is_competitive"])
        book = self.book()
        self.assertIsNone(dict(book["Сводка"].values)["Сдано"])
        self.assertTrue(
            all(row[7] == "Не требуется" for row in list(book["Решения"].values)[1:])
        )
        empty = create_partner_program(
            solution_storage_mode="solution", solution_template="hackathon"
        )
        empty.managers.add(self.manager)
        response = call(
            self.manager, "get", f"/programs/{empty.pk}/management/solution-analytics/"
        )
        self.assertEqual(response.status_code, 200)
        self.assert_schema("ManagerAnalytics", response.data)
        self.assertEqual(response.data["summary"]["solutions"]["total"], 0)
