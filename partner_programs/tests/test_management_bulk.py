from datetime import timedelta
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier, Event

from django.db import close_old_connections, connection, transaction
from django.test import TestCase, TransactionTestCase, skipUnlessDBFeature
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from partner_programs.models import (
    PartnerProgram,
    PartnerProgramProject,
    PartnerProgramUserProfile,
)
from partner_programs.tests.helpers import (
    create_partner_program,
    create_program_member,
    create_program_project,
    create_project,
    create_user,
)
from project_rates.models import (
    Criteria,
    ExpertInvitation,
    ProjectExpertAssignment,
    ProjectScore,
)
from project_rates.tests.helpers import create_rate_expert
from project_rates.workspace import revision
from projects.models import Collaborator, Project


class ParticipantRemovalConcurrencyTests(TransactionTestCase):
    @skipUnlessDBFeature("has_select_for_no_key_update")
    def test_new_score_commits_while_leader_removal_requires_fresh_preview(self):
        from partner_programs.services.participant_removal import (
            RemovalConflict,
            preview_removal,
            remove_participants,
        )

        manager = create_user()
        program = create_partner_program(is_competitive=True)
        program.managers.add(manager)
        leader = create_program_member(program)
        project = create_project(leader=leader.user)
        create_program_project(program, project=project, submitted=True)
        expert = create_rate_expert(program=program)
        criterion = Criteria.objects.create(
            partner_program=program,
            name="Результат",
            type="int",
            min_value=0,
            max_value=10,
        )
        preview = preview_removal(program.pk, manager, [leader.pk])
        program_locked, project_locked = Event(), Event()

        def score():
            close_old_connections()
            try:
                with transaction.atomic():
                    with connection.cursor() as cursor:
                        cursor.execute("SET LOCAL lock_timeout = '3s'")
                    PartnerProgram.objects.select_for_update().get(pk=program.pk)
                    program_locked.set()
                    self.assertTrue(project_locked.wait(timeout=5))
                    ProjectScore.objects.create(
                        project=project, criteria=criterion, user=expert, value="3"
                    )
                return "score committed"
            finally:
                close_old_connections()

        def remove():
            close_old_connections()

            def notice_lock(execute, sql, params, many, context):
                result = execute(sql, params, many, context)
                if 'FROM "projects_project"' in sql and " FOR " in sql:
                    project_locked.set()
                return result

            try:
                self.assertTrue(program_locked.wait(timeout=5))
                with connection.execute_wrapper(notice_lock):
                    try:
                        remove_participants(
                            program.pk, manager, [leader.pk], preview["preview_token"]
                        )
                        return "removed"
                    except RemovalConflict:
                        return "fresh preview required"
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=2) as pool:
            scoring, removal = pool.submit(score), pool.submit(remove)
            self.assertEqual(scoring.result(timeout=10), "score committed")
            self.assertEqual(removal.result(timeout=10), "fresh preview required")
        self.assertTrue(PartnerProgramUserProfile.objects.filter(pk=leader.pk).exists())
        self.assertTrue(
            ProjectScore.objects.filter(project=project, user=expert).exists()
        )

    @skipUnlessDBFeature("has_select_for_update")
    def test_same_preview_can_remove_registration_only_once(self):
        from partner_programs.services.participant_removal import (
            RemovalConflict,
            preview_removal,
            remove_participants,
        )

        manager = create_user()
        program = create_partner_program()
        program.managers.add(manager)
        member = create_program_member(program)
        preview = preview_removal(program.pk, manager, [member.pk])
        barrier = Barrier(2)

        def remove():
            close_old_connections()
            try:
                barrier.wait(timeout=5)
                try:
                    result = remove_participants(
                        program.pk, manager, [member.pk], preview["preview_token"]
                    )
                    return result["removed"]
                except RemovalConflict:
                    return "conflict"
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda _: remove(), range(2)))
        self.assertCountEqual(results, [1, "conflict"])
        self.assertFalse(PartnerProgramUserProfile.objects.filter(pk=member.pk).exists())
        self.assertTrue(member.user.__class__.objects.filter(pk=member.user_id).exists())


class ManagementBulkTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.manager = create_user()
        cls.program = create_partner_program(
            is_competitive=True, is_distributed_evaluation=True, required_evaluators=1
        )
        cls.program.managers.add(cls.manager)
        cls.foreign = create_partner_program()
        cls.leader = create_program_member(cls.program)
        cls.member = create_program_member(cls.program)
        cls.foreign_member = create_program_member(cls.foreign, user=cls.member.user)
        cls.project = create_project(leader=cls.leader.user)
        Collaborator.objects.create(project=cls.project, user=cls.member.user)
        cls.link = create_program_project(
            cls.program, project=cls.project, submitted=True
        )
        cls.foreign_link = create_program_project(cls.foreign, project=cls.project)
        cls.expert = create_rate_expert(program=cls.program)
        Criteria.objects.filter(partner_program=cls.program).delete()
        cls.criterion = Criteria.objects.create(
            partner_program=cls.program,
            name="Балл",
            type="int",
            min_value=0,
            max_value=10,
        )
        cls.assignment = ProjectExpertAssignment.objects.create(
            partner_program=cls.program, project=cls.project, expert=cls.expert.expert
        )
        cls.score = ProjectScore.objects.create(
            criteria=cls.criterion, project=cls.project, user=cls.expert, value="0"
        )
        cls.invitation = ExpertInvitation.objects.create(
            partner_program=cls.program,
            email="bulk-invite@example.test",
            token_digest="a" * 64,
            expires_at=timezone.now() + timedelta(days=2),
        )
        cls.root = f"/programs/{cls.program.pk}/management/"

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.manager)

    def post(self, suffix, data):
        return self.client.post(self.root + suffix + "/", data, format="json")

    def preview(self, ids):
        response = self.post("participants/remove-preview", {"ids": ids})
        self.assertEqual(response.status_code, 200, response.data)
        return response.data

    def test_default_twenty_and_select_all_uses_filtered_full_snapshot(self):
        for index in range(23):
            create_program_member(
                self.program, user=create_user(first_name="Выборка", city="Москва")
            )
        page = self.client.get(self.root + "participants/", {"search": "Выборка"})
        self.assertEqual(page.data["count"], 23)
        self.assertEqual(len(page.data["results"]), 20)
        selection = self.client.get(
            self.root + "participants/selection/",
            {"search": "Выборка", "limit": 1, "offset": 20},
        )
        self.assertEqual(len(selection.data["ids"]), 23)
        self.assertNotIn(self.foreign_member.pk, selection.data["ids"])
        last = self.client.get(
            self.root + "participants/", {"search": "Выборка", "offset": 20}
        )
        self.assertEqual(len(last.data["results"]), 3)

    def test_project_filter_and_selection_share_status_and_link_id_mapping(self):
        incomplete = create_program_project(self.program, submitted=True)
        not_submitted = create_program_project(self.program)
        done = self.client.get(self.root + "projects/", {"evaluation": "evaluated"})
        self.assertEqual([row["id"] for row in done.data["results"]], [self.link.pk])
        pending = self.client.get(
            self.root + "projects/", {"evaluation": "not_evaluated"}
        )
        self.assertEqual(
            {row["id"] for row in pending.data["results"]},
            {incomplete.pk, not_submitted.pk},
        )
        selection = self.client.get(
            self.root + "projects/selection/",
            {"evaluation": "evaluated", "limit": 1, "offset": 200},
        )
        self.assertEqual(selection.data["ids"], [self.link.pk])
        self.assertEqual(selection.data["projects"][0]["project_id"], self.project.pk)
        self.assertEqual(
            self.client.get(self.root + "projects/", {"evaluation": "bad"}).status_code,
            400,
        )

    def test_selected_export_preserves_activity_date_and_uses_registration_header(self):
        activity = timezone.now().replace(microsecond=0)
        type(self.member.user).objects.filter(pk=self.member.user_id).update(
            last_activity=activity
        )
        response = self.post("participants/export", {"ids": [self.member.pk]})
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(1, 5).value, "Дата регистрации (МСК)")
        self.assertEqual(sheet.cell(2, 6).data_type, "d")
        self.assertEqual(
            self.post(
                "participants/export", {"ids": [self.member.pk, self.foreign_member.pk]}
            ).status_code,
            409,
        )

    def test_member_removal_preserves_shared_team_project_account_and_other_program(self):
        preview = self.preview([self.member.pk])
        self.assertEqual(preview["applications"], [])
        response = self.post(
            "participants/remove",
            {"ids": [self.member.pk], "preview_token": preview["preview_token"]},
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertFalse(
            PartnerProgramUserProfile.objects.filter(pk=self.member.pk).exists()
        )
        self.assertTrue(
            PartnerProgramUserProfile.objects.filter(pk=self.foreign_member.pk).exists()
        )
        self.assertTrue(Project.objects.filter(pk=self.project.pk).exists())
        self.assertTrue(
            Collaborator.objects.filter(
                project=self.project, user=self.member.user
            ).exists()
        )
        self.assertTrue(
            type(self.member.user).objects.filter(pk=self.member.user_id).exists()
        )
        self.assertTrue(PartnerProgramProject.objects.filter(pk=self.link.pk).exists())

    def test_leader_preview_then_removal_unlinks_submitted_application_only(self):
        preview = self.preview([self.leader.pk])
        self.assertEqual([row["id"] for row in preview["applications"]], [self.link.pk])
        self.assertEqual(preview["assignments_count"], 1)
        response = self.post(
            "participants/remove",
            {"ids": [self.leader.pk], "preview_token": preview["preview_token"]},
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertFalse(PartnerProgramProject.objects.filter(pk=self.link.pk).exists())
        self.assertTrue(
            PartnerProgramProject.objects.filter(pk=self.foreign_link.pk).exists()
        )
        self.assertTrue(Project.objects.filter(pk=self.project.pk).exists())
        self.assertTrue(ProjectScore.objects.filter(pk=self.score.pk).exists())
        self.assertFalse(
            ProjectExpertAssignment.objects.filter(pk=self.assignment.pk).exists()
        )
        self.assertEqual(self.client.get(self.root + "projects/").data["count"], 0)

    def test_changed_rating_or_selection_requires_fresh_preview_without_partial_removal(
        self,
    ):
        preview = self.preview([self.leader.pk, self.member.pk])
        self.score.value = "7"
        self.score.save()
        response = self.post(
            "participants/remove",
            {
                "ids": [self.leader.pk, self.member.pk],
                "preview_token": preview["preview_token"],
            },
        )
        self.assertEqual(response.status_code, 409)
        self.assertTrue(
            PartnerProgramUserProfile.objects.filter(pk=self.leader.pk).exists()
        )
        self.assertTrue(
            PartnerProgramUserProfile.objects.filter(pk=self.member.pk).exists()
        )
        self.assertTrue(PartnerProgramProject.objects.filter(pk=self.link.pk).exists())
        self.assertEqual(
            self.post(
                "participants/remove",
                {"ids": [self.member.pk], "preview_token": preview["preview_token"]},
            ).status_code,
            409,
        )

    def test_preview_is_required_signed_and_program_actor_scoped(self):
        self.assertEqual(
            self.post("participants/remove", {"ids": [self.member.pk]}).status_code, 400
        )
        self.assertEqual(
            self.post(
                "participants/remove",
                {"ids": [self.member.pk], "preview_token": "forged"},
            ).status_code,
            409,
        )
        self.assertEqual(
            self.post(
                "participants/remove-preview", {"ids": [self.foreign_member.pk]}
            ).status_code,
            409,
        )
        for user, expected in [(self.member.user, 403), (None, 401)]:
            self.client.force_authenticate(user)
            for suffix in [
                "participants/remove-preview",
                "participants/remove",
                "participants/export",
                "experts/bulk",
            ]:
                self.assertEqual(
                    self.post(suffix, {"ids": [self.member.pk]}).status_code, expected
                )
            for suffix in ["participants/selection/", "projects/selection/"]:
                self.assertEqual(
                    self.client.get(self.root + suffix).status_code, expected
                )

    def test_bulk_expert_cases_and_revoke_validate_revision_and_scope(self):
        old_revision = revision(self.program)
        payload = {
            "action": "cases",
            "revision": old_revision,
            "expert_ids": [self.expert.expert.pk],
            "invitation_ids": [self.invitation.pk],
            "all_cases": True,
            "case_names": [],
        }
        response = self.post("experts/bulk", payload)
        self.assertEqual(response.status_code, 200, response.data)
        self.assertTrue(
            ProjectExpertAssignment.objects.filter(pk=self.assignment.pk).exists()
        )
        self.assertEqual(self.post("experts/bulk", payload).status_code, 409)
        self.assertEqual(
            self.post(
                "experts/bulk",
                {**payload, "action": "revoke", "revision": revision(self.program)},
            ).status_code,
            400,
        )
        response = self.post(
            "experts/bulk",
            {
                "action": "revoke",
                "revision": revision(self.program),
                "invitation_ids": [self.invitation.pk],
            },
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.invitation.refresh_from_db()
        self.assertEqual(self.invitation.status, "revoked")

    def test_bulk_resend_is_all_or_nothing_and_only_enqueues_after_commit(self):
        invitation = ExpertInvitation.objects.create(
            partner_program=self.program,
            email="second-bulk@example.test",
            token_digest="b" * 64,
            expires_at=timezone.now() + timedelta(days=1),
        )
        with self.captureOnCommitCallbacks(execute=False) as callbacks:
            response = self.post(
                "experts/bulk",
                {
                    "action": "resend",
                    "revision": revision(self.program),
                    "invitation_ids": [self.invitation.pk, invitation.pk],
                },
            )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(len(callbacks), 2)
        self.invitation.refresh_from_db()
        before = self.invitation.token_digest
        with self.captureOnCommitCallbacks(execute=False) as callbacks:
            response = self.post(
                "experts/bulk",
                {
                    "action": "resend",
                    "revision": revision(self.program),
                    "invitation_ids": [self.invitation.pk, 999999],
                },
            )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(callbacks, [])
        self.invitation.refresh_from_db()
        self.assertEqual(self.invitation.token_digest, before)
