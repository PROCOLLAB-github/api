from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from partner_programs.models import PartnerProgramField, PartnerProgramFieldValue
from partner_programs.tests.helpers import (
    create_partner_program,
    create_program_project,
    create_project,
    create_user,
)


class ManagementProjectsExportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.program = create_partner_program(is_competitive=True)
        cls.manager = create_user()
        cls.program.managers.add(cls.manager)
        cls.field = PartnerProgramField.objects.create(
            partner_program=cls.program,
            name="case",
            label="Кейс",
            field_type="select",
            options="Банк|Город",
            is_required=True,
            show_filter=True,
        )
        cls.links = []
        for i in range(25):
            project = create_project(name=f"Проект {i:02d}", draft=i == 0)
            link = create_program_project(cls.program, project=project, submitted=False)
            PartnerProgramFieldValue.objects.create(
                program_project=link,
                field=cls.field,
                value_text="Банк" if i < 23 else "Город",
            )
            type(link).objects.filter(pk=link.pk).update(submitted=i < 22)
            cls.links.append(link)
        create_program_project(
            create_partner_program(), project=create_project(name="Чужой")
        )
        cls.root = f"/programs/{cls.program.pk}/management/projects/"

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.manager)

    def test_export_matches_all_filtered_rows_before_pagination(self):
        query = {
            "search": "Проект",
            "case_scope": "selected",
            "case_name": "Банк",
            "submission": "submitted",
            "evaluation": "not_evaluated",
            "limit": 1,
            "offset": 20,
        }
        page = self.client.get(self.root, query)
        self.assertEqual(page.status_code, 200, page.data)
        self.assertEqual(page.data["count"], 22)
        response = self.client.get(self.root + "export/", query)
        self.assertEqual(response.status_code, 200)
        self.assertIn("no-store", response["Cache-Control"])
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 23)
        self.assertEqual(
            [sheet.cell(i, 1).value for i in range(2, 24)], list(range(1, 23))
        )
        self.assertEqual(sheet.cell(2, 7).value, "Черновик заявки")
        self.assertEqual(sheet.cell(2, 6).value, "Сдано")
        self.assertEqual(sheet.cell(2, 8).value, "Не оценено")

    def test_literals_empty_and_validation(self):
        project = self.links[0].project
        project.name = '=HYPERLINK("https://example.test")'
        project.save(update_fields=["name"])
        response = self.client.get(self.root + "export/", {"search": "HYPERLINK"})
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 2).value, project.name)
        self.assertEqual(sheet.cell(2, 2).data_type, "s")
        empty = self.client.get(self.root + "export/", {"search": "Не существует"})
        self.assertEqual(load_workbook(BytesIO(empty.content)).active.max_row, 1)
        for query in [
            {"evaluation": "wrong"},
            {"case_scope": "selected", "case_name": "Чужой"},
        ]:
            self.assertEqual(
                self.client.get(self.root + "export/", query).status_code, 400
            )

    def test_evaluated_filter_and_noncompetitive_export(self):
        from project_rates.models import Criteria, ProjectScore
        from project_rates.tests.helpers import create_rate_expert

        criterion = Criteria.objects.create(
            partner_program=self.program,
            name="Баллы",
            type="int",
            min_value=0,
            max_value=10,
        )
        expert = create_rate_expert(program=self.program)
        ProjectScore.objects.create(
            project=self.links[1].project, user=expert, criteria=criterion, value="7"
        )
        query = {"evaluation": "evaluated"}
        page = self.client.get(self.root, query)
        sheet = load_workbook(
            BytesIO(self.client.get(self.root + "export/", query).content)
        ).active
        self.assertEqual(page.data["count"], 1)
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 2).value, page.data["results"][0]["name"])
        self.assertEqual(sheet.cell(2, 8).value, "Оценено")
        self.program.is_competitive = False
        self.program.save(update_fields=["is_competitive"])
        sheet = load_workbook(
            BytesIO(self.client.get(self.root + "export/").content)
        ).active
        self.assertEqual(sheet.max_row, 26)
        self.assertEqual(sheet.cell(2, 6).value, "Не требуется")
        self.assertEqual(sheet.cell(2, 8).value, "Не требуется")

    def test_access_scoping(self):
        for user, status in [
            (None, 401),
            (create_user(), 403),
            (create_user(is_staff=True), 200),
        ]:
            self.client.force_authenticate(user)
            self.assertEqual(self.client.get(self.root + "export/").status_code, status)
        self.client.force_authenticate(self.manager)
        self.assertEqual(
            self.client.get("/programs/99999999/management/projects/export/").status_code,
            404,
        )
