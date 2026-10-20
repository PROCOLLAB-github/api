"""Export the same filtered queryset as the organizer's participant list."""

from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from partner_programs.services.management import person_name


def export_participants(rows):
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Участники")
    sheet.freeze_panes = "C2"
    headers = [
        "№",
        "Участник",
        "Почта",
        "Регион",
        "Дата регистрации (МСК)",
        "Последняя активность (МСК)",
    ]
    for index, width in enumerate((6, 32, 38, 28, 24, 28), 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    header_cells = []
    for title in headers:
        cell = WriteOnlyCell(sheet, title)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EEE8FB")
        header_cells.append(cell)
    sheet.append(header_cells)
    number = 0
    for number, profile in enumerate(rows.iterator(chunk_size=500), 1):
        values = [
            number,
            person_name(profile.user),
            profile.user.email,
            profile.participant_region,
            profile.datetime_created,
            profile.user.last_activity,
        ]
        cells = []
        for value in values:
            if hasattr(value, "tzinfo"):
                cell = WriteOnlyCell(
                    sheet,
                    value.astimezone(ZoneInfo("Europe/Moscow")).replace(tzinfo=None),
                )
                cell.number_format = "dd.mm.yyyy hh:mm"
            elif isinstance(value, int):
                cell = WriteOnlyCell(sheet, value)
            else:
                cell = WriteOnlyCell(
                    sheet, ILLEGAL_CHARACTERS_RE.sub("", str(value or "—"))
                )
                cell.data_type = "s"
            cells.append(cell)
        sheet.append(cells)
    sheet.auto_filter.ref = f"A1:F{number + 1}"
    answers_sheet = workbook.create_sheet("Анкеты")
    answers_sheet.append(["Участник", "Почта", "Поле анкеты", "Ответ"])
    for profile in rows.iterator(chunk_size=500):
        for field in profile.registration_schema.get("fields", []):
            value = profile.partner_program_data.get(field["id"])
            values = [
                person_name(profile.user),
                profile.user.email,
                field["label"],
                ", ".join(map(str, value)) if isinstance(value, list) else value,
            ]
            cells = []
            for item in values:
                cell = WriteOnlyCell(
                    answers_sheet,
                    ILLEGAL_CHARACTERS_RE.sub("", str(item if item is not None else "—")),
                )
                cell.data_type = "s"
                cells.append(cell)
            answers_sheet.append(cells)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
