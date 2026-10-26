"""Typed cells and real formulas, from the same results snapshot as the UI."""

from io import BytesIO
from math import isfinite

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


def text_cell(sheet, row, column, value):
    cell = sheet.cell(
        row,
        column,
        ILLEGAL_CHARACTERS_RE.sub("", str(value)) if value is not None else "",
    )
    cell.data_type = "s"  # User strings must never become executable Excel formulas.


def export_results(data):
    return _export_results(
        data, label="Проект", group_key="projects", id_key="project_id"
    )


def export_solution_results(data):
    return _export_results(
        data, label="Решение", group_key="solutions", id_key="solution_id"
    )


def _export_results(data, *, label, group_key, id_key):
    workbook = Workbook()
    workbook.calculation = CalcProperties(fullCalcOnLoad=True)
    rating = workbook.active
    rating.title = "Рейтинг"
    experts = workbook.create_sheet("Оценки экспертов")
    criteria = data["criteria"]
    numeric = [c for c in criteria if c["type"] in ("int", "float")]
    other = [c for c in criteria if c not in numeric]
    rating.append(
        [
            "Место",
            label,
            "Кейс",
            "Получено оценок",
            "Требуется оценок",
            *[c["name"] + " (среднее)" for c in numeric],
            "Итог",
            "Статус",
            "ID решения" if id_key == "solution_id" else "ID проекта",
        ]
    )
    experts.append(
        [
            label,
            "Кейс",
            "Эксперт",
            "Email",
            "Статус",
            *[c["name"] for c in numeric],
            "Сумма",
            *[c["name"] for c in other],
            "ID решения" if id_key == "solution_id" else "ID проекта",
            "ID эксперта",
        ]
    )
    total_col = 6 + len(numeric)
    for group in data["groups"]:
        for project in group[group_key]:
            expert_rows = []
            start_row = experts.max_row + 1
            for expert in project["experts"]:
                row = experts.max_row + 1
                for col, value in enumerate(
                    [
                        project["name"],
                        project["case_name"],
                        expert["name"],
                        expert["email"],
                        "Полная" if expert["complete"] else "Неполная",
                    ],
                    1,
                ):
                    text_cell(experts, row, col, value)
                values = {v["criterion_id"]: v["value"] for v in expert["values"]}
                for col, criterion in enumerate(numeric, 6):
                    value = values.get(criterion["id"])
                    if value is not None and value != "":
                        try:
                            number = float(value)
                            if not isfinite(number):
                                raise ValueError()
                            experts.cell(row, col, number)
                        except (TypeError, ValueError):
                            text_cell(experts, row, col, value)
                if expert["complete"]:
                    expert_rows.append(row)
                    if numeric:
                        experts.cell(
                            row,
                            total_col,
                            f"=SUM(F{row}:{get_column_letter(total_col - 1)}{row})",
                        )
                for col, criterion in enumerate(other, total_col + 1):
                    value = values.get(criterion["id"])
                    if criterion["type"] == "bool" and value in ("True", "False"):
                        value = "Да" if value == "True" else "Нет"
                    text_cell(experts, row, col, value)
                experts.cell(row, total_col + len(other) + 1, project[id_key])
                experts.cell(row, total_col + len(other) + 2, expert["user_id"])
            row = rating.max_row + 1
            if project["place"] is not None:
                rating.cell(row, 1, project["place"])
            text_cell(rating, row, 2, project["name"])
            text_cell(rating, row, 3, project["case_name"] or "Без кейса")
            rating.cell(row, 4, project["completed"])
            rating.cell(row, 5, project["required"])
            for index in range(len(numeric) + 1):
                if not expert_rows or not numeric:
                    continue
                col = 6 + index
                last_row = experts.max_row
                statuses = f"'Оценки экспертов'!$E${start_row}:$E${last_row}"
                column = get_column_letter(col)
                refs = f"'Оценки экспертов'!{column}{start_row}:{column}{last_row}"
                formula = f'AVERAGEIF({statuses},"Полная",{refs})'
                if index == len(numeric):
                    formula = f"ROUND({formula},2)"
                rating.cell(row, col, "=" + formula).number_format = "0.00"
            status = (
                (
                    "Оценивание завершено"
                    if project["evaluation_complete"]
                    else "Ожидает оценок"
                )
                if not numeric
                else "В рейтинге"
                if project["ready"]
                else "Предварительно"
                if project["completed"]
                else "Нет оценок"
            )
            text_cell(rating, row, total_col + 1, status)
            rating.cell(row, total_col + 2, project[id_key])
    for sheet in (rating, experts):
        sheet.freeze_panes = "D2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.data_type = "s"
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="8A63E6")
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 24
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
