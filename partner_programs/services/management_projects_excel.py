"""Export the organizer's filtered projects before pagination."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from partner_programs.services.exports import _literal_excel_cell
from partner_programs.services.management import person_name, project_data


def export_projects(rows, program):
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Проекты")
    sheet.freeze_panes = "C2"
    headers = [
        "№",
        "Проект",
        "Руководитель",
        "Почта руководителя",
        "Кейс",
        "Сдача решения",
        "Состояние заявки",
        "Оценивание",
        "Презентация",
    ]
    for index, width in enumerate((6, 38, 32, 38, 32, 20, 22, 26, 48), 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    cells = [_literal_excel_cell(sheet, title) for title in headers]
    for cell in cells:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EEE8FB")
    sheet.append(cells)
    statuses = {
        "evaluated": "Оценено",
        "partially_evaluated": "Частично оценено",
        "awaiting_evaluation": "Не оценено",
        "not_submitted": "—",
        "not_applicable": "Не требуется",
    }
    number = 0
    required = program.required_evaluators or program.max_project_rates or 1
    for number, link in enumerate(rows.iterator(chunk_size=500), 1):
        data = project_data(link, program)
        status = statuses.get(data["evaluation_status"], "—")
        if data["evaluation_status"] == "partially_evaluated":
            status += f' · {data["assignments_completed"] or 0}/{required}'
        values = [
            number,
            data["name"],
            person_name(link.project.leader),
            link.project.leader.email,
            data["case_name"] or "Без выбранного кейса",
            ("Сдано" if data["submitted"] else "Не сдано")
            if program.is_competitive
            else "Не требуется",
            "Черновик заявки" if data["draft"] else "Заявка завершена",
            status,
            data["presentation_url"] or "Не добавлена",
        ]
        sheet.append([_literal_excel_cell(sheet, value) for value in values])
    sheet.auto_filter.ref = f"A1:I{number + 1}"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
