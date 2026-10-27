"""Published native enrollment. Legacy Tilda handlers remain independent."""

import copy
import csv
import io
import math
import re
from datetime import date, timedelta
from zipfile import ZipFile, BadZipFile

from django.conf import settings
from django.core import signing
from django.db import transaction, IntegrityError
from django.shortcuts import get_object_or_404
from django.utils import timezone
from rest_framework.exceptions import ValidationError, PermissionDenied
from rest_framework_simplejwt.tokens import RefreshToken

from partner_programs.models import (
    PartnerProgram,
    PartnerProgramUserProfile,
    ProgramRegistration,
    PendingProgramRegistration,
)
from partner_programs.services.management import managed_program
from project_rates.workspace import WorkspaceConflict
from users.models import CustomUser, University

SALT = "program-native-confirm-v1"
RECEIPT_SALT = "program-native-resend-v1"
FIELD_TYPES = ("text", "textarea", "number", "date", "select", "multiselect")
SOURCES = ("own", "universities", "regions")


def frontend_url():
    return settings.EXPERT_WORKSPACE_FRONTEND_URL.rstrip("/")


def config_for(program):
    if not hasattr(program, "_native_registration_config"):
        try:
            program._native_registration_config = program.registration_config
        except ProgramRegistration.DoesNotExist:
            program._native_registration_config = None
    return program._native_registration_config


def public_link(program):
    config = config_for(program)
    if config and config.published:
        if config.published["mode"] == "native":
            return f"{frontend_url()}/program-registration/{program.pk}"
        return config.external_url if config.external_status == "connected" else None
    return program.registration_link


def initial_draft(program):
    # Do not migrate legacy JSON answers or infer field types from old arbitrary data.
    return {"mode": "external" if program.registration_link else "native", "fields": []}


def clean_schema(data, *, publishing=False):
    if not isinstance(data, dict) or set(data) != {"mode", "fields"}:
        raise ValidationError("Укажите способ регистрации и поля анкеты.")
    if data["mode"] not in ("native", "external"):
        raise ValidationError("Неизвестный способ регистрации.")
    fields = data["fields"]
    if not isinstance(fields, list) or len(fields) > 100:
        raise ValidationError("В анкете может быть не более 100 полей.")
    result, seen = [], set()
    for field in fields:
        if not isinstance(field, dict) or set(field) != {
            "id",
            "label",
            "type",
            "help_text",
            "required",
            "source",
            "options",
            "allow_other",
        }:
            raise ValidationError("Неверная конфигурация поля анкеты.")
        key = field["id"]
        if (
            not isinstance(key, str)
            or not re.fullmatch(r"f[a-f0-9]{32}", key)
            or key in seen
        ):
            raise ValidationError("У поля неверный или повторяющийся идентификатор.")
        seen.add(key)
        label, help_text = field["label"], field["help_text"]
        if not isinstance(label, str) or not label.strip() or len(label) > 255:
            raise ValidationError("Укажите название поля (до 255 символов).")
        if not isinstance(help_text, str) or len(help_text) > 1000:
            raise ValidationError("Подсказка: не более 1000 символов.")
        if field["type"] not in FIELD_TYPES or field["source"] not in SOURCES:
            raise ValidationError("Неизвестный тип поля или источник списка.")
        if not isinstance(field["required"], bool) or not isinstance(
            field["allow_other"], bool
        ):
            raise ValidationError("Некорректные настройки обязательности поля.")
        options = field["options"]
        if (
            not isinstance(options, list)
            or len(options) > 5000
            or any(
                not isinstance(v, str) or not v.strip() or len(v) > 500 for v in options
            )
        ):
            raise ValidationError(
                "В списке должно быть не более 5000 строк по 500 символов."
            )
        options = list(dict.fromkeys(v.strip() for v in options))
        if field["type"] not in ("select", "multiselect"):
            if field["source"] != "own" or options or field["allow_other"]:
                raise ValidationError(
                    "Список вариантов доступен только для выбора ответа."
                )
        elif publishing and field["source"] == "own" and not options:
            raise ValidationError(f"Добавьте варианты ответа для поля «{label}».")
        if field["source"] != "own" and options:
            raise ValidationError("Для справочника используется общий список платформы.")
        result.append(
            {
                **field,
                "label": label.strip(),
                "help_text": help_text.strip(),
                "options": options,
            }
        )
    return {"mode": data["mode"], "fields": result}


def settings_data(program):
    config = config_for(program)
    return {
        "program_draft": program.draft,
        "draft": config.draft if config else initial_draft(program),
        "published": config.published if config else None,
        "revision": config.revision if config else 0,
        "published_version": config.published_version if config else 0,
        "published_at": config.published_at if config else None,
        "link": public_link(program) if config and config.published else None,
        "external_status": config.external_status
        if config
        else ("connected" if program.registration_link else "not_connected"),
        "external_url": config.external_url
        if config
        else program.registration_link or "",
        "connection_comment": config.connection_comment if config else "",
        "connection_website_url": config.connection_website_url if config else "",
        "connection_contact": config.connection_contact if config else "",
        "answered_field_ids": list(
            {
                k
                for a in program.partner_program_profiles.values_list(
                    "partner_program_data", flat=True
                )
                if isinstance(a, dict)
                for k in a
            }
        ),
    }


def locked_config(program):
    return ProgramRegistration.objects.get_or_create(
        program=program,
        defaults={
            "draft": initial_draft(program),
            "external_status": "connected"
            if program.registration_link
            else "not_connected",
            "external_url": program.registration_link or "",
        },
    )[0]


@transaction.atomic
def save_settings(pk, user, data, *, publish=False):
    program = managed_program(pk, user, lock=True)
    config = locked_config(program)
    if data["revision"] != config.revision:
        raise WorkspaceConflict(
            "stale_registration", "Анкета изменена другим организатором. Обновите данные."
        )
    schema = clean_schema(data["draft"], publishing=publish)
    # Historical snapshot lives on each enrollment; removing an answered field archives it.
    config.draft = schema
    config.revision += 1
    if publish:
        if schema["mode"] == "external" and config.external_status != "connected":
            raise ValidationError("Дождитесь подключения внешней формы администратором.")
        if not program.draft:
            config.published = copy.deepcopy(schema)
            config.published_at = timezone.now()
            config.published_version += 1
    config.save()
    return settings_data(program)


@transaction.atomic
def request_connection(pk, user, data):
    program = managed_program(pk, user, lock=True)
    config = locked_config(program)
    if config.external_status != "connected":
        duplicate = (
            config.external_status == "requested"
            and config.connection_website_url == data["url"]
            and config.connection_comment == data["comment"]
            and config.connection_contact == data["contact"]
            and config.connection_requests.exists()
        )
        if not duplicate:
            config.connection_website_url = data["url"]
            config.connection_comment = data["comment"]
            config.connection_contact = data["contact"]
            config.requested_by = user
            config.requested_at = timezone.now()
            config.external_status = "requested"
            config.revision += 1
            config.save()
            from partner_programs.services.registration_notifications import (
                record_request,
            )

            record_request(config, user)
    return settings_data(program)


def open_program(pk, *, lock=False):
    rows = PartnerProgram.objects.select_for_update() if lock else PartnerProgram.objects
    return get_object_or_404(rows, pk=pk, draft=False)


def published_config(program):
    config = config_for(program)
    if not config or not config.published or config.published["mode"] != "native":
        raise ValidationError("Встроенная регистрация еще не опубликована.")
    return config


def check_open(program):
    if program.datetime_registration_ends < timezone.now():
        raise WorkspaceConflict("registration_closed", "Регистрация завершена.")


def public_data(program, user):
    config = published_config(program)
    enrolled = bool(
        user.is_authenticated
        and PartnerProgramUserProfile.objects.filter(
            partner_program=program, user=user
        ).exists()
    )
    return {
        "id": program.pk,
        "name": program.name,
        "version": config.published_version,
        "fields": config.published["fields"],
        "closed": program.datetime_registration_ends < timezone.now(),
        "enrolled": enrolled,
        "authenticated": user.is_authenticated,
        "program_url": f"/office/program/{program.pk}",
        "registration_ends": program.datetime_registration_ends,
    }


def validate_answers(schema, answers):
    if not isinstance(answers, dict):
        raise ValidationError({"answers": "Укажите ответы анкеты."})
    fields = {f["id"]: f for f in schema["fields"]}
    if set(answers) - set(fields):
        raise ValidationError(
            {"answers": "В форме есть неизвестные поля. Обновите страницу."}
        )
    errors, clean = {}, {}
    for key, field in fields.items():
        value = answers.get(key)
        if value is None or value == "" or value == []:
            if field["required"]:
                errors[key] = "Заполните обязательное поле."
            continue
        kind = field["type"]
        valid = True
        if kind in ("text", "textarea", "date"):
            valid = (
                isinstance(value, str)
                and bool(value.strip())
                and len(value) <= (10000 if kind == "textarea" else 1000)
            )
            if valid and kind == "date":
                try:
                    date.fromisoformat(value)
                except ValueError:
                    valid = False
        elif kind == "number":
            valid = (
                not isinstance(value, bool)
                and isinstance(value, (int, float))
                and math.isfinite(value)
                and abs(value) < 1e15
            )
        else:
            values = value if kind == "multiselect" else [value]
            valid = (
                isinstance(values, list)
                and len(values) <= 100
                and all(isinstance(v, str) and 0 < len(v) <= 500 for v in values)
            )
            if valid:
                if field["source"] == "own":
                    options = field["options"]
                elif field["source"] == "universities":
                    options = list(
                        University.objects.filter(
                            name__in=values, is_active=True
                        ).values_list("name", flat=True)
                    )
                else:
                    from partner_programs.registration_regions import REGIONS

                    options = REGIONS
                valid = all(
                    v in options
                    or (
                        field["allow_other"]
                        and v.startswith("Другое: ")
                        and bool(v[8:].strip())
                    )
                    for v in values
                )
        if valid:
            clean[key] = value.strip() if isinstance(value, str) else value
        else:
            errors[key] = "Укажите корректное значение из формы."
    if errors:
        raise ValidationError({"answers": errors})
    return clean


def current_answers(program, data):
    check_open(program)
    config = published_config(program)
    if config.published_version != data["version"]:
        raise WorkspaceConflict(
            "form_changed", "Анкета обновлена. Проверьте ответы и отправьте ее еще раз."
        )
    return config, validate_answers(config.published, data["answers"])


def create_enrollment(program, user, answers, schema):
    region = next(
        (
            answers.get(f["id"], "")
            for f in schema.get("fields", [])
            if f["source"] == "regions" and f["type"] == "select"
        ),
        "",
    )
    profile, created = PartnerProgramUserProfile.objects.get_or_create(
        partner_program=program,
        user=user,
        defaults={
            "partner_program_data": answers,
            "registration_schema": schema,
            "registration_region": region
            if isinstance(region, str) and not region.startswith("Другое: ")
            else "",
        },
    )
    return profile


@transaction.atomic
def enroll(pk, user, data):
    if not user.is_authenticated or not user.is_active:
        raise PermissionDenied("Войдите в подтвержденный аккаунт.")
    program = open_program(pk, lock=True)
    if PartnerProgramUserProfile.objects.filter(
        partner_program=program, user=user
    ).exists():
        return {"enrolled": True, "program_id": pk}
    config, answers = current_answers(program, data)
    create_enrollment(program, user, answers, config.published)
    return {"enrolled": True, "program_id": pk}


def confirmation_token(pending):
    return signing.dumps(
        {"id": str(pending.pk), "generation": pending.generation}, salt=SALT
    )


def receipt_token(pending):
    return signing.dumps({"id": str(pending.pk)}, salt=RECEIPT_SALT)


def queue_confirmation(pending):
    from partner_programs.tasks import send_program_registration_confirmation

    def dispatch():
        try:
            send_program_registration_confirmation.delay(
                str(pending.pk), pending.generation
            )
        except Exception:
            PendingProgramRegistration.objects.filter(pk=pending.pk).update(
                delivery_status="failed"
            )

    transaction.on_commit(dispatch)


@transaction.atomic
def signup(pk, data):
    program = open_program(pk, lock=True)
    config, answers = current_answers(program, data)
    email = data["email"].strip().lower()
    if CustomUser.objects.filter(email__iexact=email).exists():
        raise WorkspaceConflict(
            "account_exists",
            "Аккаунт с этой почтой уже есть. Войдите, чтобы участвовать.",
        )
    try:
        with transaction.atomic():
            user = CustomUser.objects.create_user(
                email=email,
                password=data["password"],
                first_name=data["first_name"],
                last_name=data["last_name"],
                birthday=data["birthday"],
                is_active=False,
                onboarding_stage=None,
            )
            pending = PendingProgramRegistration.objects.create(
                program=program, user=user, answers=answers, schema=config.published
            )
    except IntegrityError:
        raise WorkspaceConflict(
            "account_exists", "Аккаунт уже создан. Войдите, чтобы участвовать."
        )
    queue_confirmation(pending)
    return {"email_pending": True, "receipt": receipt_token(pending), "email": email}


@transaction.atomic
def confirm(pk, token):
    try:
        proof = signing.loads(token, salt=SALT, max_age=86400)
    except signing.BadSignature:
        raise WorkspaceConflict(
            "confirmation_expired",
            "Ссылка недействительна или истекла. Отправьте письмо еще раз.",
        )
    program = open_program(pk, lock=True)
    pending = get_object_or_404(
        PendingProgramRegistration.objects.select_for_update().select_related("user"),
        pk=proof.get("id"),
        program=program,
    )
    if pending.confirmed_at or pending.generation != proof.get("generation"):
        raise WorkspaceConflict(
            "confirmation_used", "Ссылка уже использована. Войдите в аккаунт."
        )
    user = pending.user
    user.is_active = True
    user.save(update_fields=["is_active"])
    config = config_for(program)
    can_enroll = bool(
        config
        and config.published
        and config.published["mode"] == "native"
        and program.datetime_registration_ends >= timezone.now()
    )
    review_required = False
    if PartnerProgramUserProfile.objects.filter(
        partner_program=program, user=user
    ).exists():
        can_enroll = True
    elif can_enroll:
        current_keys = {f["id"] for f in config.published["fields"]}
        try:
            answers = validate_answers(
                config.published,
                {k: v for k, v in pending.answers.items() if k in current_keys},
            )
        except ValidationError:
            can_enroll, review_required = False, True
        else:
            create_enrollment(program, user, answers, config.published)
    pending.confirmed_at = timezone.now()
    pending.save(update_fields=["confirmed_at"])
    refresh = RefreshToken.for_user(user)
    return {
        "access": str(refresh.access_token),
        "refresh": str(refresh),
        "program_id": pk,
        "enrolled": can_enroll,
        "review_required": review_required,
        "answers": pending.answers if review_required else {},
    }


@transaction.atomic
def resend(pk, token):
    try:
        proof = signing.loads(token, salt=RECEIPT_SALT, max_age=7 * 86400)
    except signing.BadSignature:
        try:
            proof = signing.loads(token, salt=SALT, max_age=7 * 86400)
        except signing.BadSignature:
            raise ValidationError(
                "Ссылка недействительна. Используйте вход или восстановление пароля."
            )
    program = open_program(pk, lock=True)
    pending = get_object_or_404(
        PendingProgramRegistration.objects.select_for_update(),
        pk=proof.get("id"),
        program=program,
    )
    if pending.confirmed_at:
        raise WorkspaceConflict(
            "confirmation_used", "Почта уже подтверждена. Войдите в аккаунт."
        )
    if pending.last_sent_at > timezone.now() - timedelta(seconds=60):
        raise WorkspaceConflict(
            "resend_wait", "Повторное письмо можно отправить через минуту."
        )
    pending.generation += 1
    pending.last_sent_at = timezone.now()
    pending.delivery_status = "queued"
    pending.save(update_fields=["generation", "last_sent_at", "delivery_status"])
    queue_confirmation(pending)
    return {"email_pending": True, "receipt": receipt_token(pending)}


def import_options(upload):
    if upload.size > 2 * 1024 * 1024:
        raise ValidationError("Размер файла должен быть не больше 2 МБ.")
    raw = upload.read()
    try:
        if upload.name.lower().endswith(".csv"):
            content = raw.decode("utf-8-sig")
            dialect = (
                csv.Sniffer().sniff(content[:4096], delimiters=",;\t")
                if any(c in content for c in ",;\t")
                else csv.excel
            )
            rows = csv.reader(io.StringIO(content), dialect)
        elif upload.name.lower().endswith(".xlsx"):
            with ZipFile(io.BytesIO(raw)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError("Archive too large")
            from openpyxl import load_workbook

            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            rows = book.active.iter_rows(values_only=True)
        else:
            raise ValueError("Unsupported file")
        options, seen, duplicates, blanks = [], set(), 0, 0
        for index, row in enumerate(rows):
            if index >= 10000 or len(row) > 1:
                raise ValueError("One column and at most 10000 rows required")
            value = str(row[0]).strip() if row and row[0] is not None else ""
            if len(value) > 500 or value.startswith("="):
                raise ValueError("Invalid option")
            if not value:
                blanks += 1
            elif value.casefold() in seen:
                duplicates += 1
            else:
                options.append(value)
                seen.add(value.casefold())
            if len(options) > 5000:
                raise ValueError("Too many options")
        if not options:
            raise ValueError("Empty list")
        return {"options": options, "duplicates": duplicates, "blanks": blanks}
    except (
        ValueError,
        UnicodeError,
        BadZipFile,
        csv.Error,
        KeyError,
        OSError,
        TypeError,
    ):
        raise ValidationError(
            "Не удалось прочитать файл. Нужен XLSX/CSV: один столбец без заголовка, до 5000 вариантов."
        )
