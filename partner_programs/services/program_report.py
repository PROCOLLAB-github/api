"""A single, program-scoped workbook with stable IDs between its sheets."""

from collections import Counter
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import json
from zoneinfo import ZoneInfo

from django.db.models import Prefetch
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from partner_programs.models import PartnerProgramFieldValue, PartnerProgramProject
from partner_programs.services.project_analytics import (
    build_project_analytics as build_program_manager_analytics,
)
from partner_programs.services.project_assignment_analytics import build_assignments
from partner_programs.services.project_case_analytics import project_case_links
from partner_programs.services.case_fields import get_program_case_field
from partner_programs.services.exports import _literal_excel_cell
from partner_programs.services.analytics_actions_support import (
    participant_rows,
    person_name,
)
from project_rates.models import Criteria, ProjectScore
from projects.models import Collaborator

ASSIGNMENT_LABELS = {
    "completed": "Выполнено",
    "pending": "Ожидает оценки",
    "in_progress": "В работе",
    "not_ready": "Решение еще не сдано",
}


def text_value(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def score_value(value, kind):
    if value is not None and value != "" and kind in ("int", "float"):
        try:
            number = Decimal(str(value))
            if number.is_finite():
                return number
        except InvalidOperation:
            pass
    return value


class Report:
    def __init__(self):
        self.book = Workbook(write_only=True)

    def sheet(self, name, headers):
        sheet = self.book.create_sheet(name)
        sheet.freeze_panes = "A2"
        for index, title in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(index)].width = (
                30 if "ID" not in title else 16
            )
        cells = []
        for title in headers:
            cell = _literal_excel_cell(sheet, title)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E8E0FA")
            cells.append(cell)
        sheet.append(cells)
        return sheet

    def append(self, sheet, values):
        cells = []
        for value in values:
            if isinstance(value, datetime):
                value = value.astimezone(ZoneInfo("Europe/Moscow")).replace(tzinfo=None)
                cell = WriteOnlyCell(sheet, value)
                cell.number_format = "dd.mm.yyyy hh:mm"
            elif isinstance(value, Decimal):
                cell = WriteOnlyCell(sheet, value)
            else:
                cell = _literal_excel_cell(sheet, text_value(value))
            cells.append(cell)
        sheet.append(cells)

    def finish(self):
        output = BytesIO()
        self.book.save(output)
        return output.getvalue()


def build_program_report(program):
    report = Report()
    overview = build_program_manager_analytics(program)
    summary = report.sheet("Сводка", ["Показатель", "Значение"])
    for row in [
        ["ID программы", program.pk],
        ["Программа", program.name],
        ["Сформировано (МСК)", timezone.now()],
        ["Описание", program.description],
        ["Начало программы (МСК)", program.datetime_started],
        ["Окончание программы (МСК)", program.datetime_finished],
        ["Окончание регистрации (МСК)", program.datetime_registration_ends],
        ["Окончание приема решений (МСК)", program.get_project_submission_deadline()],
        ["Окончание оценивания (МСК)", program.datetime_evaluation_ends],
        ["Участники", overview["summary"]["participants"]["total"]],
        ["Проекты", overview["summary"]["projects"]["total"]],
        ["Сдано", overview["solution_funnel"]["submitted"]],
        ["Оценено проектов", overview["solution_funnel"]["evaluated"]],
        ["Заходили на платформу", overview["participant_funnel"]["visited_participants"]],
        ["Вступили в команду", overview["participant_funnel"]["with_team"]],
        [
            "Формат оценивания",
            "Распределенное" if program.is_distributed_evaluation else "Открытое",
        ],
        [
            "Связи между листами",
            "ID участника = ID пользователя. Проекты, команды, назначения и оценки связаны по ID проекта.",
        ],
    ]:
        report.append(summary, row)

    people = report.sheet(
        "Участники",
        [
            "ID участника",
            "Участник",
            "Почта",
            "Регион",
            "Дата регистрации (МСК)",
            "Последняя активность (МСК)",
            "Последний вход (МСК)",
        ],
    )
    answers = report.sheet(
        "Анкеты",
        ["ID регистрации", "ID участника", "Участник", "ID поля", "Поле анкеты", "Ответ"],
    )
    seen = set()
    for profile in participant_rows(program).iterator(chunk_size=500):
        user = profile.user
        if user.pk not in seen:
            report.append(
                people,
                [
                    user.pk,
                    person_name(user),
                    user.email,
                    profile.participant_region,
                    profile.datetime_created,
                    user.last_activity,
                    user.last_login,
                ],
            )
            seen.add(user.pk)
        schema = (
            getattr(profile, "registration_schema", program.data_schema)
            if isinstance(
                getattr(profile, "registration_schema", program.data_schema), dict
            )
            else {}
        )
        fields = {
            field["id"]: field.get("label", field["id"])
            for field in schema.get("fields", [])
            if isinstance(field, dict) and "id" in field
        }
        data = (
            profile.partner_program_data
            if isinstance(profile.partner_program_data, dict)
            else {}
        )
        for key, value in data.items():
            report.append(
                answers,
                [
                    profile.pk,
                    user.pk,
                    person_name(user),
                    key,
                    fields.get(key, key),
                    value,
                ],
            )

    fields = list(program.fields.order_by("pk"))
    projects = report.sheet(
        "Проекты",
        [
            "ID проекта",
            "ID связи с программой",
            "Проект",
            "Описание",
            "Кейс",
            "Регион",
            "ID лидера",
            "Лидер",
            "Размер команды",
            "Презентация",
            "Сдача решения",
            "Добавлен (МСК)",
            "Сдан (МСК)",
        ]
        + [f"{field.label} [ID {field.pk}]" for field in fields],
    )
    teams = report.sheet(
        "Команды", ["ID проекта", "Проект", "ID участника", "Участник", "Роль"]
    )
    case_field = get_program_case_field(program)
    links = (
        project_case_links(program, field=case_field)
        .select_related("project", "project__leader")
        .prefetch_related(
            Prefetch(
                "project__collaborator_set",
                queryset=Collaborator.objects.select_related("user"),
                to_attr="report_collaborators",
            ),
            Prefetch(
                "field_values",
                queryset=PartnerProgramFieldValue.objects.filter(
                    field__partner_program=program
                ),
                to_attr="report_values",
            ),
        )
        .order_by("pk")
    )
    for link in links.iterator(chunk_size=500):
        project = link.project
        team = {}
        if project.leader_id:
            team[project.leader_id] = (project.leader, "Лидер")
        for member in project.report_collaborators:
            team.setdefault(member.user_id, (member.user, "Участник"))
        values = {value.field_id: value.value_text for value in link.report_values}
        report.append(
            projects,
            [
                project.pk,
                link.pk,
                project.name,
                project.description,
                link.case_name or "Без выбранного кейса",
                project.region,
                project.leader_id,
                person_name(project.leader) if project.leader_id else "",
                len(team),
                project.presentation_address,
                ("Сдано" if link.submitted else "Не сдано")
                if program.is_competitive
                else "Не требуется",
                link.datetime_created,
                link.datetime_submitted if link.submitted else None,
            ]
            + [values.get(field.pk, "") for field in fields],
        )
        for user, role in team.values():
            report.append(
                teams, [project.pk, project.name, user.pk, person_name(user), role]
            )

    cases = report.sheet("Кейсы", ["Кейс", "Проектов", "Сдано", "Участников"])
    for case in overview["cases"]["items"]:
        report.append(
            cases,
            [
                case["name"],
                case["projects_total"],
                case["submitted"],
                case["participants_total"],
            ],
        )
    case = overview["cases"]["without_case"]
    report.append(
        cases,
        [
            "Без выбранного кейса",
            case["projects_total"],
            case["submitted"],
            case["participants_total"],
        ],
    )

    assignment_data = build_assignments(program.pk)
    total = Counter(row["expert"]["expert_id"] for row in assignment_data)
    completed = Counter(
        row["expert"]["expert_id"]
        for row in assignment_data
        if row["status"] == "completed"
    )
    experts = report.sheet(
        "Эксперты",
        ["ID эксперта", "ID пользователя", "Эксперт", "Почта", "Назначений", "Выполнено"],
    )
    for expert in program.experts.select_related("user").order_by("pk"):
        report.append(
            experts,
            [
                expert.pk,
                expert.user_id,
                person_name(expert.user),
                expert.user.email,
                total[expert.pk],
                completed[expert.pk],
            ],
        )
    assignments = report.sheet(
        "Назначения",
        [
            "ID назначения",
            "ID проекта",
            "Проект",
            "ID эксперта",
            "ID пользователя",
            "Эксперт",
            "Статус",
            "Назначено (МСК)",
            "Сдано (МСК)",
        ],
    )
    for row in assignment_data:
        report.append(
            assignments,
            [
                row["assignment_id"],
                row["project"]["id"],
                row["project"]["name"],
                row["expert"]["expert_id"],
                row["expert"]["user_id"],
                row["expert"]["full_name"],
                ASSIGNMENT_LABELS[row["status"]],
                row["assigned_at"],
                row["project_submitted_at"],
            ],
        )
    criteria = report.sheet(
        "Критерии", ["ID критерия", "Критерий", "Описание", "Тип", "Минимум", "Максимум"]
    )
    for criterion in Criteria.objects.filter(partner_program=program).order_by("pk"):
        report.append(
            criteria,
            [
                criterion.pk,
                criterion.name,
                criterion.description,
                criterion.type,
                criterion.min_value,
                criterion.max_value,
            ],
        )
    scores = report.sheet(
        "Оценки",
        [
            "ID оценки",
            "ID проекта",
            "Проект",
            "ID пользователя эксперта",
            "Эксперт",
            "ID критерия",
            "Критерий",
            "Тип",
            "Значение",
            "Исходное значение",
        ],
    )
    for score in (
        ProjectScore.objects.filter(
            criteria__partner_program=program,
            project_id__in=PartnerProgramProject.objects.filter(
                partner_program=program
            ).values("project_id"),
        )
        .select_related("criteria", "project", "user")
        .order_by("project_id", "user_id", "criteria_id")
    ):
        report.append(
            scores,
            [
                score.pk,
                score.project_id,
                score.project.name,
                score.user_id,
                person_name(score.user),
                score.criteria_id,
                score.criteria.name,
                score.criteria.type,
                score_value(score.value, score.criteria.type),
                score.value,
            ],
        )
    return report.finish()
