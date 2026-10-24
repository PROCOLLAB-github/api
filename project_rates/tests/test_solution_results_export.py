from io import BytesIO
from urllib.parse import urlencode

from django.test import TestCase, override_settings
from openpyxl import load_workbook

from partner_programs.services.results_excel import export_solution_results
from partner_programs.tests.helpers import create_user
from partner_programs.tests.test_solution_forms import call
from project_rates.models import Criteria
from project_rates.solution_models import SolutionScore
from project_rates.solution_results import results_data
from project_rates.tests.helpers import create_rate_program
from project_rates.tests.test_solution_evaluation import EvaluationFixture


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class SolutionResultsExportTests(EvaluationFixture, TestCase):
    def setUp(self):
        self.seed_evaluation()
        self.export_url = (
            f"/programs/{self.program.pk}/management/solution-results/export/"
        )
        self.results_url = f"/programs/{self.program.pk}/management/solution-results/"

    def download(self, **query):
        response = call(self.manager, "get", self.export_url + "?" + urlencode(query))
        self.assertEqual(response.status_code, 200, getattr(response, "data", None))
        self.assertEqual(response["Cache-Control"], "private, no-store")
        self.assertEqual(response["Referrer-Policy"], "no-referrer")
        return response, load_workbook(BytesIO(response.content), data_only=False)

    def test_real_ids_json_parity_numeric_zero_false_and_formula_types(self):
        self.rate()
        response, book = self.download()
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(
            f'filename="program-{self.program.pk}-solution-results.xlsx"',
            response["Content-Disposition"],
        )
        self.assertEqual(book.sheetnames, ["Рейтинг", "Оценки экспертов"])
        rating, experts = book.worksheets
        payload = call(self.manager, "get", self.results_url).data
        row = payload["groups"][0]["solutions"][0]
        self.assertEqual(rating["B1"].value, "Решение")
        self.assertEqual(rating["I1"].value, "ID решения")
        self.assertEqual(rating["A2"].value, row["place"])
        self.assertEqual(rating["B2"].value, row["name"])
        self.assertEqual(rating["D2"].value, row["completed"])
        self.assertEqual(rating["E2"].value, row["required"])
        self.assertEqual(rating["I2"].value, self.solution.pk)
        self.assertEqual(experts["J1"].value, "ID решения")
        self.assertEqual(experts["J2"].value, self.solution.pk)
        self.assertEqual(experts["K2"].value, self.expert.pk)
        self.assertEqual(experts["F2"].value, 0)
        self.assertEqual(experts["F2"].data_type, "n")
        self.assertEqual(experts["H2"].value, "Нет")
        self.assertEqual(experts["G2"].value, "=SUM(F2:F2)")
        self.assertEqual(experts["G2"].data_type, "f")
        self.assertEqual(
            rating["G2"].value,
            "=ROUND(AVERAGEIF('Оценки экспертов'!$E$2:$E$2,\"Полная\","
            "'Оценки экспертов'!G2:G2),2)",
        )
        self.assertTrue(book.calculation.fullCalcOnLoad)
        self.assertEqual(row["score"], 0)
        self.assertNotIn("ID проекта", [c.value for c in rating[1]])

    def test_all_scope_filters_and_drafts_have_same_selection_as_results(self):
        self.rate()
        second = self.another_solution()
        draft = self.another_solution(submitted=False)
        _, all_book = self.download(scope="all", search="missing", status="pending")
        self.assertEqual(
            {row[8] for row in all_book["Рейтинг"].values if isinstance(row[8], int)},
            {self.solution.pk, second.pk},
        )
        self.assertNotIn(draft.pk, [cell.value for cell in all_book["Рейтинг"]["I"][1:]])
        _, missing = self.download(search="missing")
        self.assertEqual(missing["Рейтинг"].max_row, 1)
        _, filtered = self.download(search="Вторая", status="pending")
        self.assertEqual(filtered["Рейтинг"].max_row, 2)
        self.assertEqual(filtered["Рейтинг"]["I2"].value, second.pk)
        _, complete = self.download(
            status="complete", case_scope="selected", case_name="Задача 1"
        )
        self.assertEqual(complete["Рейтинг"].max_row, 2)
        self.assertEqual(complete["Рейтинг"]["I2"].value, self.solution.pk)
        _, without = self.download(case_scope="without_case")
        self.assertEqual(without["Рейтинг"].max_row, 1)

    def test_partial_expert_remains_in_book_but_not_in_average_formula(self):
        self.rate()
        SolutionScore.objects.create(
            solution=self.solution,
            user=self.second,
            criteria=self.criteria[0],
            value="10",
        )
        _, book = self.download()
        experts = book["Оценки экспертов"]
        self.assertEqual(experts["E2"].value, "Полная")
        self.assertEqual(experts["E3"].value, "Неполная")
        self.assertIsNone(experts["G3"].value)
        self.assertIn("AVERAGEIF(", book["Рейтинг"]["G2"].value)
        self.assertIn('"Полная"', book["Рейтинг"]["G2"].value)
        self.assertEqual(book["Рейтинг"]["D2"].value, 1)

    def test_user_text_never_becomes_formula_and_writer_does_not_mutate_data(self):
        self.rate()
        query = dict(search="", case_scope="all", scope="all", order="desc", status="all")
        data = results_data(self.program, query)
        row = data["groups"][0]["solutions"][0]
        row["name"] = '=HYPERLINK("https://example.test","text")'
        row["experts"][0]["values"][-1]["value"] = "=1+1"
        book = load_workbook(BytesIO(export_solution_results(data)))
        self.assertEqual(book["Рейтинг"]["B2"].value, row["name"])
        self.assertEqual(book["Рейтинг"]["B2"].data_type, "s")
        self.assertEqual(book["Оценки экспертов"]["I2"].value, "=1+1")
        self.assertEqual(book["Оценки экспертов"]["I2"].data_type, "s")
        self.assertNotIn("project_id", row)
        self.assertNotIn("projects", data["groups"][0])

    def test_empty_and_non_numeric_results_are_valid_without_invented_place(self):
        _, empty = self.download(search="missing")
        self.assertEqual([sheet.max_row for sheet in empty.worksheets], [1, 1])
        Criteria.objects.filter(pk=self.criteria[0].pk).delete()
        self.rate(scores=[dict(criterion_id=self.criteria[1].pk, value="False")])
        _, book = self.download()
        self.assertIsNone(book["Рейтинг"]["A2"].value)
        self.assertIsNone(book["Рейтинг"]["F2"].value)
        self.assertEqual(book["Рейтинг"]["G2"].value, "Оценивание завершено")

    def test_roles_other_program_and_legacy_mode_cannot_download_closed_results(self):
        for actor in (None, self.captain, self.expert, create_user()):
            with self.subTest(actor=getattr(actor, "pk", None)):
                response = call(actor, "get", self.export_url)
                self.assertIn(response.status_code, (401, 403))
                self.assert_schema("Error", response.data)
        foreign = create_rate_program()
        response = call(
            self.manager,
            "get",
            f"/programs/{foreign.pk}/management/solution-results/export/",
        )
        self.assertEqual(response.status_code, 403)
        foreign.managers.add(self.manager)
        response = call(
            self.manager,
            "get",
            f"/programs/{foreign.pk}/management/solution-results/export/",
        )
        self.assertEqual(response.status_code, 409)
        self.assert_schema("Error", response.data)
        response = call(
            self.manager,
            "get",
            "/programs/9007199254740991/management/solution-results/export/",
        )
        self.assertEqual(response.status_code, 404)

    def test_invalid_query_uses_c00_json_errors_including_unknown_keys(self):
        for query in (
            dict(scope="wrong"),
            dict(project_id=1),
            dict(case_scope="selected", case_name="Чужой кейс"),
            dict(scope="all", case_scope="selected", case_name="Чужой кейс"),
        ):
            with self.subTest(query=query):
                response = call(
                    self.manager, "get", self.export_url + "?" + urlencode(query)
                )
                self.assertEqual(response.status_code, 400, response.data)
                self.assert_schema("Error", response.data)
                self.assertIn("application/json", response["Content-Type"])
